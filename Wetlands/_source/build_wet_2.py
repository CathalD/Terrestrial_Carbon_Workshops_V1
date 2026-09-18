"""Stage 2 — Plot & Site Log, Core Log, Peat Data, Chronology."""
import sys, json; sys.path.insert(0, "/home/user/Terrestrial_Carbon_Workshops_V1/Wetlands/_source")
from wet_style import *
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

SCR = "/tmp/claude-0/-home-user/af4fc136-c459-5930-bae9-67144880d936/scratchpad/wet"
R = json.load(open(f"{SCR}/_wref.json"))
wb = openpyxl.load_workbook(f"{SCR}/_w1.xlsx")
N_PLOT, N_CORE, N_PEAT, N_CHRON = 60, 120, 600, 200
HR = 4                      # header row on every data sheet
F1 = HR + 1                 # first data row

def build(title, pos, hdr, widths, fills, nrows, formulas, dvs=None, note=None, freeze="C"):
    ws = wb.create_sheet(title, pos)
    ws["A1"] = title.split(". ", 1)[1].upper()
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    if note:
        ws["A2"] = note; ws["A2"].font = F_IT
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(hdr))
        ws["A2"].alignment = WRAP; ws.row_dimensions[2].height = 34
    header_row(ws, HR, hdr, widths=widths)
    for i, f in enumerate(fills, start=1):
        if f: paint(ws, F1, F1 + nrows - 1, i, f)
    for col, t in formulas.items():
        for r in range(F1, F1 + nrows):
            ws.cell(r, col, t.format(r=r))
    for rng, dv in (dvs or []):
        d = DataValidation(type="list", formula1=dv, allow_blank=True, showDropDown=False)
        ws.add_data_validation(d); d.add(rng.format(a=F1, b=F1 + nrows - 1))
    ws.freeze_panes = f"{freeze}{F1}"
    return ws

RF = "'R1. Reference'"
VP_RNG = f"{RF}!$A${R['VP_FIRST']}:$A${R['VP_LAST']}"
WT_RNG = f"{RF}!$A${R['WT_FIRST']}:$A${R['WT_LAST']}"

# ═════════════════════════════════════════ 1. PLOT & SITE LOG
hdr = ["Plot ID", "Site ID", "Study area", "Date", "Location", "Latitude", "Longitude",
       "GNSS accuracy (m)", "Datum", "Elevation (m)", "Wetland type", "Microform",
       "Vegetation zone", "Water table depth (cm)", "Plot area (m²)", "Notes"]
w  = [13, 10, 12, 11, 20, 11, 11, 12, 10, 11, 13, 13, 16, 13, 12, 32]
fl = ["y"] * 15 + ["y"]
ws1 = build("1. Plot & Site Log", 1, hdr, w, fl, N_PLOT, {},
  dvs=[("K{a}:K{b}", f"={WT_RNG}"),
       ("L{a}:L{b}", '"Hummock,Hollow,Lawn,Carpet,Not recorded"'),
       ("O{a}:O{b}", '"1,4,16,25,100,400"')],
  note="ONE ROW PER PLOT. Fill this in first — everything joins to it on Plot ID. Water table depth is "
       "measured DOWN from the peat surface, so a positive number means the water table sits below the "
       "surface and a negative number means standing water above it. Microform matters: hollows accumulate "
       "more organic matter than hummocks, so recording it lets you stratify later.")
PL = "'1. Plot & Site Log'"
PKEY = f"{PL}!$A${F1}:$A${F1+N_PLOT-1}"
PSITE = f"{PL}!$B${F1}:$B${F1+N_PLOT-1}"
PAREA = f"{PL}!$O${F1}:$O${F1+N_PLOT-1}"

# ═════════════════════════════════════════ 2. CORE LOG
PD = "'3. Peat Data'"
pd_core = f"{PD}!$B${F1}:$B${F1+N_PEAT-1}"
pd_bot  = f"{PD}!$F${F1}:$F${F1+N_PEAT-1}"
pd_full = f"{PD}!$P${F1}:$P${F1+N_PEAT-1}"
pd_ref  = f"{PD}!$Q${F1}:$Q${F1+N_PEAT-1}"
hdr = ["Core ID", "Plot ID", "Latitude", "Longitude", "Coring date", "Drives taken",
       "Bore depth (cm)", "Core length recovered (cm)", "Surface compaction (cm)",
       "Recovery ratio", "Depth to mineral contact (cm)", "Reached mineral contact?",
       "von Post at base", "Chamber length (cm)", "Chamber diameter (cm)", "Photo series?",
       "Deepest section logged (cm)", "Stock, full profile (kg C/m²)",
       "Stock to reference depth (kg C/m²)", "Reaches reference depth?", "QC flags", "Notes"]
w  = [13, 13, 11, 11, 11, 10, 12, 14, 13, 11, 14, 14, 12, 13, 13, 11, 14, 16, 16, 14, 46, 26]
fl = ["y"] * 9 + ["g"] + ["y", "y", "y", "y", "y", "y"] + ["g"] * 5 + ["y"]
f = {
 10: '=IF(OR($G{r}="",$H{r}="",$G{r}=0),"",$H{r}/$G{r})',
 17: (f'=IF($A{{r}}="","",IF(SUMPRODUCT(MAX(({pd_core}=$A{{r}})*{pd_bot}))=0,"",'
      f'SUMPRODUCT(MAX(({pd_core}=$A{{r}})*{pd_bot}))))'),
 18: f'=IF($A{{r}}="","",IFERROR(SUMIFS({pd_full},{pd_core},$A{{r}}),""))',
 19: f'=IF($A{{r}}="","",IFERROR(SUMIFS({pd_ref},{pd_core},$A{{r}}),""))',
 20: '=IF(NOT(ISNUMBER($Q{r})),"",IF($Q{r}>=REFERENCE_DEPTH_CM,"Yes","No"))',
 21: ('=IF($A{r}="","",'
      f'IF(AND($B{{r}}<>"",COUNTIF({PKEY},$B{{r}})=0),"Plot ID not in Plot & Site Log. ","")'
      '&IF(NOT(ISNUMBER($Q{r})),"No peat sections logged for this core. ","")'
      '&IF($T{r}="No","Core stops short of the reference depth, so its reference-depth stock is an underestimate. Its full-profile stock is still valid. ","")'
      '&IF($L{r}="No","Core did not reach the mineral contact — the full-profile stock is a MINIMUM, not the whole store. ","")'
      '&IF(AND(ISNUMBER($K{r}),$K{r}<PEATLAND_MIN_DEPTH_CM),"Organic depth under "&PEATLAND_MIN_DEPTH_CM&" cm: by the guide\'s definition this is not peatland. Report it, but say so. ","")'
      '&IF(AND(ISNUMBER($J{r}),$J{r}<QC_RECOVERY_MIN),"Recovered only "&TEXT($J{r},"0%")&" of the bore depth — the chamber probably did not fill completely. Material is missing, which is not the same as compaction. ","")'
      '&IF(AND(ISNUMBER($J{r}),$J{r}>1.05),"Recovered more than the bore depth — check the two measurements. ","")'
      '&IF(AND(ISNUMBER($I{r}),$I{r}>2),"Surface compaction of "&$I{r}&" cm: depths in this core sit shallower than in the ground, so depth-limited stock is biased. Full-profile stock is unaffected. ",""))'),
}
ws2 = build("2. Core Log", 2, hdr, w, fl, N_CORE, f,
  dvs=[("B{a}:B{b}", f"={PKEY}"), ("L{a}:L{b}", '"Yes,No,Unsure"'),
       ("P{a}:P{b}", '"Yes,No"'), ("M{a}:M{b}", f"={VP_RNG}")],
  note="ONE ROW PER CORE. A Russian (Macaulay) corer takes 50 cm at a time, so a 2 m core is four drives — "
       "log the core once here and its sections on the next sheet. BORE DEPTH is how deep the corer went; "
       "CORE LENGTH RECOVERED is what came back. For a side-filling corer these should be nearly equal: a "
       "shortfall means the chamber did not fill, NOT that the peat compressed.")
CL = "'2. Core Log'"
CKEY = f"{CL}!$A${F1}:$A${F1+N_CORE-1}"

# ═════════════════════════════════════════ 3. PEAT DATA
hdr = ["Plot ID", "Core ID", "Drive #", "Section ID", "Top depth (cm)", "Bottom depth (cm)",
       "Thickness (cm)", "von Post", "Bulk density basis", "Bulk density (g/cm³)",
       "Water content (%)", "LOI₅₅₀ (%)", "Organic carbon (%)", "Carbon % used",
       "Thickness within reference depth (cm)", "Stock, full thickness (kg C/m²)",
       "Stock to reference depth (kg C/m²)", "_stock per cm",
       "Cumulative carbon to base (kg C/m²)", "QC flags", "Notes"]
w  = [12, 13, 8, 10, 11, 12, 11, 10, 19, 12, 11, 11, 12, 11, 15, 15, 15, 13, 16, 46, 26]
fl = ["y"] * 9 + ["b", "b", "b", "b"] + ["g"] * 7 + ["y"]
f = {
 7:  '=IF(OR($E{r}="",$F{r}=""),"",$F{r}-$E{r})',
 14: '=IF($M{r}<>"",$M{r},IF($L{r}<>"",$L{r}*CARBON_FRACTION_OM,""))',
 15: '=IF(OR($E{r}="",$F{r}=""),"",MAX(0,MIN($F{r},REFERENCE_DEPTH_CM)-MIN($E{r},REFERENCE_DEPTH_CM)))',
 16: '=IF(OR($J{r}="",$N{r}="",$G{r}=""),"",$J{r}*($N{r}/100)*$G{r}*10)',
 17: '=IF(OR($J{r}="",$N{r}="",$O{r}=""),"",$J{r}*($N{r}/100)*$O{r}*10)',
 18: '=IF(OR(NOT(ISNUMBER($P{r})),NOT(ISNUMBER($G{r})),$G{r}=0),0,$P{r}/$G{r})',  # 0 not "" — SUMPRODUCT on sheet 4 multiplies this range
 19: (f'=IF(OR($B{{r}}="",NOT(ISNUMBER($F{{r}}))),"",'
      f'SUMIFS({pd_full},{pd_core},$B{{r}},{pd_bot},"<="&$F{{r}}))'),
 20: ('=IF(AND($A{r}="",$B{r}="",$E{r}=""),"",'
      f'IF(AND($B{{r}}<>"",COUNTIF({CKEY},$B{{r}})=0),"Core ID not in Core Log. ","")'
      f'&IF(AND($A{{r}}<>"",COUNTIF({PKEY},$A{{r}})=0),"Plot ID not in Plot & Site Log. ","")'
      '&IF(AND(ISNUMBER($E{r}),ISNUMBER($F{r}),$F{r}<=$E{r}),"Bottom depth is not below top depth. ","")'
      '&IF($I{r}="","Say which bulk-density basis the lab reported. ","")'
      '&IF(AND(ISNUMBER($J{r}),OR($J{r}<QC_BD_MIN,$J{r}>QC_BD_MAX)),"Bulk density outside "&QC_BD_MIN&"–"&QC_BD_MAX&" g/cm³ for peat. Above the top of that range usually means the mineral contact has been passed. ","")'
      '&IF(AND(ISNUMBER($N{r}),$N{r}>QC_CARBON_PCT_MAX),"Carbon above "&QC_CARBON_PCT_MAX&"% — check units; this may be organic matter rather than carbon. ","")'
      '&IF(AND($L{r}<>"",$M{r}<>""),"Both LOI and measured carbon given — the measured value is used. ","")'
      '&IF(AND($M{r}="",$L{r}<>"",OM_FACTOR_IS_LOCAL<>"Yes"),"Carbon derived from LOI using the default factor of "&CARBON_FRACTION_OM&", which has not been calibrated locally. See R1. Reference. ","")'
      '&IF(AND(ISNUMBER($O{r}),$O{r}=0,ISNUMBER($E{r})),"Section lies entirely below the reference depth — it counts to the full-profile stock only. ",""))'),
}
ws3 = build("3. Peat Data", 3, hdr, w, fl, N_PEAT, f,
  dvs=[("A{a}:A{b}", f"={PKEY}"), ("B{a}:B{b}", f"={CKEY}"), ("H{a}:H{b}", f"={VP_RNG}"),
       ("I{a}:I{b}", '"Whole sample / sample volume,Fine fraction / fine-fraction volume"')],
  note="ONE ROW PER SECTION. Depths run continuously down the core, so the bottom of one section is the top "
       "of the next — and they continue across drive boundaries (drive 2 starts where drive 1 ended). Enter "
       "either LOI₅₅₀ or a measured organic carbon %; if you have both, the measured value wins.")
ws3.column_dimensions["R"].hidden = True

# ═════════════════════════════════════════ 4. CHRONOLOGY
CH = "'4. Chronology'"
ch_core = f"{CH}!$A${F1}:$A${F1+N_CHRON-1}"
ch_dep  = f"{CH}!$B${F1}:$B${F1+N_CHRON-1}"
ch_age  = f"{CH}!$C${F1}:$C${F1+N_CHRON-1}"
ch_cum  = f"{CH}!$F${F1}:$F${F1+N_CHRON-1}"
ch_rec  = f"{CH}!$J${F1}:$J${F1+N_CHRON-1}"
pd_spcm = f"{PD}!$R${F1}:$R${F1+N_PEAT-1}"
pd_top  = f"{PD}!$E${F1}:$E${F1+N_PEAT-1}"
hdr = ["Core ID", "Depth (cm)", "Age (yr before coring)", "Age ± (yr)", "Date source",
       "Cumulative carbon at depth (kg C/m²)", "Interval SAR (cm/yr)",
       "Interval CAR (g C/m²/yr)", "QC flags", "_recent", "_prevdepth"]
w  = [13, 11, 15, 10, 22, 17, 13, 15, 44, 9, 10]
fl = ["y", "y", "y", "y", "y"] + ["g"] * 4 + ["g", "g"]
f = {
 # cumulative carbon at an arbitrary depth = whole sections above it + the part-section it falls in
 6:  (f'=IF(OR($A{{r}}="",NOT(ISNUMBER($B{{r}}))),"",'
      f'SUMIFS({pd_full},{pd_core},$A{{r}},{pd_bot},"<="&$B{{r}})'
      f'+SUMPRODUCT(({pd_core}=$A{{r}})*({pd_top}<$B{{r}})*({pd_bot}>$B{{r}})*{pd_spcm}*($B{{r}}-{pd_top})))'),
 10: '=IF($E{r}="",0,IF(OR(LEFT($E{r},6)="Pb-210",LEFT($E{r},6)="Cs-137"),1,0))',
 11: (f'=IF(OR($A{{r}}="",NOT(ISNUMBER($B{{r}}))),"",'
      f'SUMPRODUCT(MAX(({ch_core}=$A{{r}})*({ch_dep}<$B{{r}})*{ch_dep})))'),
 7:  (f'=IF(OR($A{{r}}="",NOT(ISNUMBER($B{{r}})),NOT(ISNUMBER($C{{r}}))),"",'
      f'IFERROR(($B{{r}}-$K{{r}})/($C{{r}}-IF($K{{r}}=0,0,SUMIFS({ch_age},{ch_core},$A{{r}},{ch_dep},$K{{r}}))),""))'),
 8:  (f'=IF(OR($A{{r}}="",NOT(ISNUMBER($F{{r}})),NOT(ISNUMBER($C{{r}}))),"",'
      f'IFERROR((($F{{r}}-IF($K{{r}}=0,0,SUMIFS({ch_cum},{ch_core},$A{{r}},{ch_dep},$K{{r}})))*1000)'
      f'/($C{{r}}-IF($K{{r}}=0,0,SUMIFS({ch_age},{ch_core},$A{{r}},{ch_dep},$K{{r}}))),""))'),
 9:  ('=IF($A{r}="","",'
      f'IF(COUNTIF({CKEY},$A{{r}})=0,"Core ID not in Core Log. ","")'
      '&IF($E{r}="","No date source given — RERCA and LORCA need to know which method produced this age. ","")'
      '&IF(AND(ISNUMBER($C{r}),$C{r}<0),"Negative age: this horizon post-dates the coring year. Check the conversion. ","")'
      '&IF(AND(ISNUMBER($G{r}),$G{r}<0),"Age reversal — this horizon is younger than one above it. A mixed or disturbed profile may not be datable at all. ","")'
      '&IF(AND(ISNUMBER($D{r}),ISNUMBER($C{r}),$C{r}>0,$D{r}/$C{r}>0.5),"Age uncertainty exceeds 50% of the age. Usable as an indication only. ","")'
      '&IF(AND($D{r}="",$C{r}<>""),"No uncertainty given. Every radiometric age has one; a rate without it cannot be defended. ",""))'),
}
ws4 = build("4. Chronology", 4, hdr, w, fl, N_CHRON, f,
  dvs=[("A{a}:A{b}", f"={CKEY}"),
       ("E{a}:E{b}", '"Pb-210 CRS,Pb-210 Plum,Pb-210 CFCS,Cs-137 (1963 peak),Cs-137 (1954 onset),'
                     'C-14 calibrated,Stratigraphic marker,Other"')],
  note="OPTIONAL — only if the core has been dated. ONE ROW PER DATED HORIZON. Age is in years BEFORE THE "
       "CORING YEAR, not years BP (which is referenced to 1950). The age-depth modelling itself happens in R "
       "— see Part 5 — and what you type here are its outputs. Carbon comes from the peat sheet, so nothing "
       "is entered twice.")
for c in ("J", "K"): ws4.column_dimensions[c].hidden = True

wb.save(f"{SCR}/_w2.xlsx")
json.dump(dict(F1=F1, N_PLOT=N_PLOT, N_CORE=N_CORE, N_PEAT=N_PEAT, N_CHRON=N_CHRON),
          open(f"{SCR}/_wrows.json", "w"))
print("stage 2 ok:", wb.sheetnames)
