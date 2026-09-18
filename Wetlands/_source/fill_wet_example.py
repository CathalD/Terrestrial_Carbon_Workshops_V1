"""Fill the wetland calculator with a constructed teaching dataset."""
import openpyxl, random
SRC = "/home/user/Terrestrial_Carbon_Workshops_V1/Wetlands/04_Data_Interpretation/calculators/Wetland_Carbon_Calculator.xlsx"
OUT = "/tmp/claude-0/-home-user/af4fc136-c459-5930-bae9-67144880d936/scratchpad/wet/_filled.xlsx"
random.seed(20260918)
wb = openpyxl.load_workbook(SRC)
F1 = 5

# ── 1. Plot & Site Log ───────────────────────────────────────────────────────
ws = wb["1. Plot & Site Log"]
plots = [
 # plot, site, area, date, location, lat, lon, gnss, datum, elev, type, microform, zone, WT, area
 ("MB-01","S1","MB","2026-08-04","Mica Bog centre",   52.1140,-83.4420,0.4,"WGS84",196,"Bog","Hummock","Sphagnum lawn",  12,100),
 ("MB-02","S1","MB","2026-08-04","Mica Bog centre",   52.1152,-83.4398,0.5,"WGS84",196,"Bog","Hollow", "Sphagnum lawn",   4,100),
 ("MB-03","S1","MB","2026-08-05","Mica Bog margin",   52.1178,-83.4361,0.4,"WGS84",194,"Bog","Lawn",   "Lagg margin",     9,100),
 ("RF-01","S2","MB","2026-08-06","Rushing Fen",       52.1042,-83.4512,0.5,"WGS84",189,"Fen","Lawn",   "Sedge lawn",      2,100),
 ("RF-02","S2","MB","2026-08-06","Rushing Fen",       52.1031,-83.4488,0.6,"WGS84",189,"Fen","Hollow", "Sedge lawn",     -3,100),
 ("RF-03","S2","MB","2026-08-07","Rushing Fen edge",  52.1058,-83.4535,0.5,"WGS84",190,"Fen","Hummock","Shrub fen",       8,100),
 ("CS-01","S3","MB","2026-08-08","Cedar Swamp",       52.0975,-83.4610,0.7,"WGS84",185,"Swamp","Hummock","Treed swamp",   6,100),
 ("CS-02","S3","MB","2026-08-08","Cedar Swamp",       52.0962,-83.4592,0.6,"WGS84",185,"Swamp","Hollow", "Treed swamp",   1,100),
 ("CS-03","S3","MB","2026-08-09","Cedar Swamp edge",  52.0988,-83.4634,0.8,"WGS84",186,"Swamp","Lawn",   "Swamp margin", 14,100),
]
for i, p in enumerate(plots):
    for j, v in enumerate(p, start=1): ws.cell(F1+i, j, v)

# ── profiles ─────────────────────────────────────────────────────────────────
# (plot, core, total peat depth, reached mineral, recovery quirk, fine-sectioned?)
cores = [
 ("MB-01","MB-01-C1", 352, "Yes", 1.00, True),   # the dated core
 ("MB-01","MB-01-C2", 338, "Yes", 1.00, False),
 ("MB-02","MB-02-C1", 366, "Yes", 1.00, False),
 ("MB-03","MB-03-C1", 214, "Yes", 0.86, False),  # poor recovery — chamber did not fill
 ("RF-01","RF-01-C1", 182, "Yes", 1.00, False),
 ("RF-02","RF-02-C1", 196, "Yes", 1.00, False),
 ("RF-03","RF-03-C1", 158, "Yes", 1.00, False),
 ("CS-01","CS-01-C1",  96, "Yes", 1.00, False),
 ("CS-02","CS-02-C1", 104, "No",  1.00, False),  # refusal on buried wood, contact not reached
 ("CS-03","CS-03-C1",  26, "Yes", 1.00, False),  # under 30 cm — not peatland by definition
]
TYPE = {p[0]: p[10] for p in plots}

def props(wtype, mid, total):
    """Bulk density and LOI at a mid-section depth, by wetland type."""
    f = mid / max(total, 1)
    if wtype == "Bog":
        bd  = 0.035 + 0.115 * f ** 0.7 + random.uniform(-0.006, 0.008)
        loi = 98.5 - 7.0 * f + random.uniform(-1.2, 0.8)
    elif wtype == "Fen":
        bd  = 0.085 + 0.130 * f ** 0.7 + random.uniform(-0.010, 0.012)
        loi = 91.0 - 13.0 * f + random.uniform(-2.5, 1.8)
    else:  # Swamp
        bd  = 0.140 + 0.190 * f ** 0.7 + random.uniform(-0.015, 0.018)
        loi = 82.0 - 22.0 * f + random.uniform(-4.0, 3.0)
    return round(max(0.025, bd), 3), round(min(99.0, max(30.0, loi)), 1)

cl = wb["2. Core Log"]; pd_ = wb["3. Peat Data"]; cr = sr = F1
for plot, cid, depth, reached, recov, fine in cores:
    wtype = TYPE[plot]
    drives = -(-depth // 50)
    bore = depth
    length = round(depth * recov, 1)
    for j, v in enumerate([cid, plot, round(52.10+random.uniform(-.01,.01),5),
                           round(-83.45+random.uniform(-.01,.01),5), "2026-08-08", drives,
                           bore, length, round(random.uniform(0.0, 1.4), 1)], start=1):
        cl.cell(cr, j, v)
    cl.cell(cr, 11, depth if reached == "Yes" else "")
    cl.cell(cr, 12, reached)
    cl.cell(cr, 13, "H7" if wtype == "Swamp" else ("H5" if wtype == "Fen" else "H4"))
    cl.cell(cr, 14, 50); cl.cell(cr, 15, 5.2); cl.cell(cr, 16, "Yes")
    if reached == "No":
        cl.cell(cr, 22, "Refusal on buried wood at 104 cm; mineral contact not reached.")
    cr += 1
    # section edges
    edges = []
    if fine:
        edges += list(range(0, 30, 2)) + list(range(30, 50, 5))
        start = 50
    else:
        start = 0
    e = start
    while e < depth:
        edges.append(e); e += 10
    edges.append(depth)
    edges = sorted(set(edges))
    for si in range(len(edges) - 1):
        top, bot = edges[si], edges[si+1]
        bd, loi = props(wtype, (top+bot)/2, depth)
        drive = int(top // 50) + 1
        vp = "H3" if (top+bot)/2 < depth*0.25 else ("H5" if (top+bot)/2 < depth*0.7 else "H7")
        for j, v in enumerate([plot, cid, drive, si+1, top, bot], start=1):
            pd_.cell(sr, j, v)
        pd_.cell(sr, 8, vp)
        pd_.cell(sr, 9, "Whole sample / sample volume")
        pd_.cell(sr, 10, bd)
        pd_.cell(sr, 11, round(100 - bd*55 + random.uniform(-3, 3), 1))
        pd_.cell(sr, 12, loi)
        sr += 1
    # the mineral contact section, where the signature flips
    if reached == "Yes":
        bd_m = round(random.uniform(0.85, 1.15), 3)
        for j, v in enumerate([plot, cid, drives, len(edges), depth, depth + 6], start=1):
            pd_.cell(sr, j, v)
        pd_.cell(sr, 8, "H10"); pd_.cell(sr, 9, "Whole sample / sample volume")
        pd_.cell(sr, 10, bd_m); pd_.cell(sr, 11, round(random.uniform(25, 40), 1))
        pd_.cell(sr, 12, round(random.uniform(4, 11), 1))
        pd_.cell(sr, 21, "Mineral contact — grey silty clay.")
        sr += 1
N_PEAT_ROWS = sr - F1

# ── 4. Chronology: one fully dated bog core ─────────────────────────────────
ch = wb["4. Chronology"]
dated = [
 # depth, age (yr before coring), +-, source
 ( 6,   14,  4, "Pb-210 CRS"),
 (12,   33,  7, "Pb-210 CRS"),
 (18,   63, 11, "Cs-137 (1963 peak)"),
 (24,   96, 16, "Pb-210 CRS"),
 (30,  127, 24, "Pb-210 CRS"),
 (75,  430, 60, "C-14 calibrated"),
 (150, 1710,110, "C-14 calibrated"),
 (240, 3990,150, "C-14 calibrated"),
 (352, 6820,190, "C-14 calibrated"),
]
for i, (d, a, u, s) in enumerate(dated):
    for j, v in enumerate(["MB-01-C1", d, a, u, s], start=1):
        ch.cell(F1+i, j, v)

# ── 5. Plot Summary: vegetation carbon carried in from the Forests calculator ─
ps = wb["5. Plot Summary"]
for i, p in enumerate(plots):
    if p[10] == "Swamp":
        ps.cell(5+i, 6, round(random.uniform(3.4, 6.1), 2))   # treed swamp: real tree carbon
    else:
        ps.cell(5+i, 6, round(random.uniform(0.15, 0.55), 2)) # shrubs and ground layer only

# ── 6. Site Summary ─────────────────────────────────────────────────────────
ss = wb["6. Site Summary"]
for i, (sid, area) in enumerate([("S1", 240000), ("S2", 155000), ("S3", 88000)]):
    ss.cell(5+i, 1, sid); ss.cell(5+i, 2, area)

wb.save(OUT)
print(f"cores={len(cores)} peat_sections={N_PEAT_ROWS} dated_horizons={len(dated)} -> {OUT}")
