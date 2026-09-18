#!/usr/bin/env python3
"""
Build the two Wetlands .docx field assets:

  03_Field_Methods/datasheets/Peat-Core-Data-Sheet.docx
  03_Field_Methods/checklists/Peat_Carbon_Skill_Checklist.docx

Both are modelled on the Forests equivalents so the series looks like one thing:
  Forests/03_Field_Methods/datasheets/Soil-Carbon-Data-Sheet.docx
  Forests/03_Field_Methods/checklists/Soil_Carbon_Skill_Checklist.docx

The palette is lifted from the Forests checklist: WWF orange e86825 for section
bands, f5dbc4 for sub-section bands, fff5ee/ffffff alternating data rows,
d9d9d9 for the column header.

Field names match the calculator tabs exactly (1. Plot & Site Log, 2. Core Log,
3. Peat Data) so transcription is a straight copy, which is what Part 3 claims.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
WET = os.path.dirname(HERE)

ORANGE   = "e86825"   # WWF orange — section bands
PALE     = "f5dbc4"   # sub-section bands
TINT     = "fff5ee"   # alternating data row
WHITE    = "ffffff"
GREY     = "d9d9d9"   # column header
FONT     = "Calibri"

# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #

def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexfill)
    tcPr.append(shd)


def borders(table, sz=4, nil=False):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    el = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'nil' if nil else 'single')
        e.set(qn('w:sz'), '0' if nil else str(sz))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), '000000')
        el.append(e)
    tblPr.append(el)


def no_autofit(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(old)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    table.autofit = False


def row_height(table, inches, first=0):
    """Minimum row height, so there is room to actually write in the box."""
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    for row in table.rows[first:]:
        row.height = Inches(inches)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def set_widths(table, widths):
    """Fix column widths for real.

    python-docx only writes per-cell w:tcW. Word mostly honours that, but
    LibreOffice sizes fixed-layout tables from w:tblGrid, which python-docx
    leaves at its autofit defaults -- so without this the columns come out
    roughly equal regardless of what you asked for. Write both.
    """
    no_autofit(table)
    tbl = table._tbl
    for old in tbl.findall(qn('w:tblGrid')):
        tbl.remove(old)
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(int(w.twips)))
        grid.append(gc)
    # tblGrid must sit immediately after tblPr
    tbl.insert(list(tbl).index(tbl.tblPr) + 1, grid)
    for row in table.rows:
        row_cells = row.cells
        for i, w in enumerate(widths):
            if i < len(row_cells):
                row_cells[i].width = w


def write(cell, text, bold=False, size=9, color=None, align=None, italic=False):
    """Replace a cell's content with a single run."""
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = FONT
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def write_lines(cell, lines, size=9, bold_first=False):
    """Multi-line cell: list of strings, or (text, bold) tuples."""
    cell.text = ''
    first = True
    for item in lines:
        text, bold = (item, False) if isinstance(item, str) else item
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = bold or (bold_first and first)
        r.font.size = Pt(size)
        r.font.name = FONT
        first = False


def para(doc, text, bold=False, size=10, align=None, space_after=4, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = FONT
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def band(table, row_idx, text, fill, color=None, size=9, bold=True):
    """Merge a row across and shade it as a section band."""
    row = table.rows[row_idx]
    merged = row.cells[0]
    for c in row.cells[1:]:
        merged = merged.merge(c)
    write(merged, text, bold=bold, size=size, color=color)
    shade(merged, fill)
    return merged


def set_margins(section, inches):
    section.left_margin = Inches(inches)
    section.right_margin = Inches(inches)
    section.top_margin = Inches(inches)
    section.bottom_margin = Inches(inches)


def base_doc(margin=0.75):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10)
    set_margins(doc.sections[0], margin)
    return doc


def landscape(section):
    w, h = section.page_width, section.page_height
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = max(w, h), min(w, h)


# --------------------------------------------------------------------------- #
# von Post scale — shared by both documents
# --------------------------------------------------------------------------- #

VON_POST = [
    ("H1",  "Completely undecomposed",      "Clear, colourless water. Plant remains easily identifiable."),
    ("H2",  "Almost undecomposed",          "Clear, yellowish water. Plant remains still identifiable."),
    ("H3",  "Very weakly decomposed",       "Slightly turbid brown water. NO peat escapes between the fingers."),
    ("H4",  "Weakly decomposed",            "Strongly turbid brown water. No peat escapes. Residue mushy."),
    ("H5",  "Moderately decomposed",        "Very turbid water, a LITTLE peat escaping. Structure evident but indistinct."),
    ("H6",  "Fairly well decomposed",       "About ONE THIRD escapes between the fingers. Residue very mushy."),
    ("H7",  "Strongly decomposed",          "About HALF escapes. Plant structure barely discernible."),
    ("H8",  "Very strongly decomposed",     "About TWO THIRDS escapes. Only resistant remains — roots, fibres."),
    ("H9",  "Almost completely decomposed", "Nearly all escapes as a fairly uniform paste."),
    ("H10", "Completely decomposed",        "ALL escapes between the fingers. No free water visible."),
]


def von_post_table(doc, size=8):
    t = doc.add_table(rows=1 + len(VON_POST), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    set_widths(t, [Inches(0.6), Inches(1.9), Inches(4.6)])
    for j, hdr in enumerate(("Class", "Degree of decomposition",
                             "What you see when you squeeze a handful of WET peat")):
        write(t.rows[0].cells[j], hdr, bold=True, size=size)
        shade(t.rows[0].cells[j], GREY)
    for i, (cls, deg, obs) in enumerate(VON_POST, start=1):
        fill = TINT if i % 2 else WHITE
        for j, v in enumerate((cls, deg, obs)):
            write(t.rows[i].cells[j], v, bold=(j == 0), size=size)
            shade(t.rows[i].cells[j], fill)
    return t


# --------------------------------------------------------------------------- #
# 1. Peat Core Data Sheet
# --------------------------------------------------------------------------- #

PLOT_FIELDS = [
    ("Plot ID:", "Site ID:", "Study area:"),
    ("Location:", "Date:", "Recorder:"),
    ("Latitude:", "Longitude:", "GNSS accuracy (m):"),
    ("Datum (e.g. WGS84):", "Elevation (m):", "Plot area (m²):"),
]

PLOT_PEAT_FIELDS = [
    ("Wetland type  (bog / fen / swamp / marsh / other):", None, None),
    ("Microform at the coring point  (hummock / hollow / lawn):", None, None),
    ("Vegetation zone  (e.g. Sphagnum lawn, sedge lawn, shrub fen, treed swamp, lagg margin):", None, None),
    ("Water table depth (cm)   — measured DOWN from the peat surface;  + = below surface,  − = standing water above:", None, None),
    ("Weather / conditions:", None, None),
]

CORE_FIELDS_LEFT = [
    "Core ID:",
    "Coring date:",
    "Latitude:",
    "Longitude:",
    "Drives taken (50 cm each):",
    "Bore depth (cm):",
    "Core length recovered (cm):",
]
CORE_FIELDS_RIGHT = [
    "Surface compaction (cm):",
    "Depth to mineral contact (cm):",
    "Reached mineral contact?        YES  ☐      NO  ☐",
    "von Post at base (H1–H10):",
    "Chamber length (cm):",
    "Chamber diameter (cm):",
    "Photo series taken (14)?        YES  ☐      NO  ☐",
]


def build_datasheet(path):
    doc = base_doc(margin=0.7)

    para(doc, "WWF-Canada Carbon Training", size=9, color="7F7F7F", space_after=0)
    para(doc, "Peat Core Data Sheet", bold=True, size=16, space_after=2)
    para(doc, "Field sheet for peatland soil carbon — bogs, fens and swamps. "
              "Field names match the Wetland Carbon Calculator tabs 1. Plot & Site Log, "
              "2. Core Log and 3. Peat Data, so transcription is a straight copy.",
         size=8, italic=True, space_after=8)

    para(doc, "Project name:  ______________________________________________     "
              "Page ______ of ______", size=10, space_after=8)

    # ---- Plot & Site Log ------------------------------------------------- #
    t = doc.add_table(rows=1 + len(PLOT_FIELDS) + 1 + len(PLOT_PEAT_FIELDS), cols=3)
    borders(t)
    set_widths(t, [Inches(2.4), Inches(2.4), Inches(2.3)])

    band(t, 0, "1.  PLOT & SITE LOG      — one block per plot; fill this in FIRST",
         ORANGE, color="FFFFFF", size=10)

    r = 1
    for triple in PLOT_FIELDS:
        for j, label in enumerate(triple):
            write(t.rows[r].cells[j], (label or "") + "\n", size=9)
            shade(t.rows[r].cells[j], WHITE)
        r += 1

    band(t, r, "Peatland-specific — none of these are optional",
         PALE, size=9)
    r += 1
    for label, _, _ in PLOT_PEAT_FIELDS:
        merged = t.rows[r].cells[0].merge(t.rows[r].cells[1]).merge(t.rows[r].cells[2])
        write(merged, label + "   ______________________________________", size=9)
        shade(merged, WHITE)
        r += 1

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    for r in sp.runs:
        r.font.size = Pt(6)

    # ---- Core Log -------------------------------------------------------- #
    t2 = doc.add_table(rows=2 + max(len(CORE_FIELDS_LEFT), len(CORE_FIELDS_RIGHT)) + 2, cols=2)
    borders(t2)
    set_widths(t2, [Inches(3.55), Inches(3.55)])

    band(t2, 0, "2.  CORE LOG      — one block per core", ORANGE, color="FFFFFF", size=10)
    hint = t2.rows[1].cells[0].merge(t2.rows[1].cells[1])
    write(hint, "Recovery ratio = core length recovered ÷ bore depth.  Below 0.90 the chamber "
                "probably did not fill — material is MISSING, which is not the same as compaction. "
                "Do not apply a compaction correction.", size=8, italic=True)
    shade(hint, PALE)

    for i in range(max(len(CORE_FIELDS_LEFT), len(CORE_FIELDS_RIGHT))):
        row = t2.rows[2 + i]
        left = CORE_FIELDS_LEFT[i] if i < len(CORE_FIELDS_LEFT) else ""
        right = CORE_FIELDS_RIGHT[i] if i < len(CORE_FIELDS_RIGHT) else ""
        for j, label in enumerate((left, right)):
            write(row.cells[j], label + "\n", size=9)
            shade(row.cells[j], WHITE)

    nrow = 2 + max(len(CORE_FIELDS_LEFT), len(CORE_FIELDS_RIGHT))
    m = t2.rows[nrow].cells[0].merge(t2.rows[nrow].cells[1])
    write_lines(m, [("Reason for refusal  (buried wood / rock / gravel / rods bending / reached contact):", True),
                    ""], size=9)
    shade(m, WHITE)
    m2 = t2.rows[nrow + 1].cells[0].merge(t2.rows[nrow + 1].cells[1])
    write_lines(m2, [("Core notes — colour and texture changes, gradual or distinct, wood, gaps, "
                      "water saturation, drive boundaries:", True), "", "", ""], size=9)
    shade(m2, WHITE)

    # ---- Page 2: peat sections ------------------------------------------ #
    doc.add_page_break()
    para(doc, "3.  PEAT DATA — core sections", bold=True, size=13, space_after=2)
    para(doc, "Core ID:  ____________________        Sections cut at:  1 cm  ☐    2 cm  ☐    "
              "5 cm  ☐    at distinct layers  ☐", size=9, space_after=2)
    para(doc, "Bulk density basis the lab reports:   whole sample \u2610      "
              "fine fraction (<2 mm) \u2610      not yet confirmed \u2610", size=9, space_after=2)
    para(doc, "Record TOP depth and BOTTOM depth, not thickness \u2014 the calculator derives "
              "thickness and checks for gaps and overlaps. Bulk density, water content, LOI and "
              "organic carbon % come back from the lab and are typed straight into the "
              "calculator; they are not field measurements. Note the DRIVE each section came "
              "from \u2014 sections straddling a drive boundary are the ones most likely to be "
              "wrong.", size=8, italic=True, space_after=6)
    NSEC = 28
    t3 = doc.add_table(rows=1 + NSEC, cols=6)
    borders(t3)
    widths = [Inches(0.7), Inches(0.55), Inches(0.8), Inches(0.85), Inches(0.7), Inches(3.5)]
    set_widths(t3, widths)
    heads = ["Section ID", "Drive #", "Top depth (cm)", "Bottom depth (cm)", "von Post",
             "Colour · texture · wood · gaps · saturation · notes"]
    for j, (h, w) in enumerate(zip(heads, widths)):
        write(t3.rows[0].cells[j], h, bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t3.rows[0].cells[j], GREY)
    for i in range(1, NSEC + 1):
        fill = TINT if i % 2 else WHITE
        for j in range(len(widths)):
            write(t3.rows[i].cells[j], "", size=9)
            shade(t3.rows[i].cells[j], fill)
    row_height(t3, 0.26, first=1)

    # ---- Page 3: depth survey grid -------------------------------------- #
    doc.add_page_break()
    para(doc, "Within-plot peat depth survey — 10 × 10 m plot, 1 m grid", bold=True, size=13, space_after=2)
    para(doc, "Plot ID:  ____________________        Probe depths in cm.        "
              "121 points.", size=10, space_after=2)
    para(doc, "Choose the coring point at the MEDIAN of these depths, not the mean and not the "
              "deepest — one very deep reading is usually a hole between hummocks or a probe that "
              "punched into soft mineral. Keep this page: pooled across plots these depths are the "
              "variability prior for your next campaign.", size=8, italic=True, space_after=6)

    t4 = doc.add_table(rows=12, cols=12)
    borders(t4)
    set_widths(t4, [Inches(0.6)] + [Inches(0.56)] * 11)
    write(t4.rows[0].cells[0], "m ↓ / →", bold=True, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(t4.rows[0].cells[0], GREY)
    for j in range(1, 12):
        write(t4.rows[0].cells[j], str(j - 1), bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t4.rows[0].cells[j], GREY)
    for i in range(1, 12):
        write(t4.rows[i].cells[0], str(i - 1), bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t4.rows[i].cells[0], GREY)
        for j in range(1, 12):
            write(t4.rows[i].cells[j], "", size=9)
            shade(t4.rows[i].cells[j], WHITE)
    row_height(t4, 0.32, first=1)

    doc.add_paragraph()
    para(doc, "Median depth:  ________ cm         Coring point chosen at:  "
              "x = ____ m,  y = ____ m,  depth ________ cm", size=10, space_after=2)
    para(doc, "If the coring point is NOT at the median depth, say why "
              "(tree / root mass / surface log / water too deep / ground unstable):",
         size=9, space_after=2)
    para(doc, "_______________________________________________________________________________"
              "____________________", size=10, space_after=10)
    para(doc, "Probe depth is not core depth. A probe can stop on buried wood or a dense layer "
              "well above the true mineral contact, and can punch through soft mineral and "
              "overstate depth. Record both — Bore depth and Depth to mineral contact are "
              "separate fields on page 1.", size=8, italic=True)

    # ---- Page 4: von Post reference ------------------------------------- #
    doc.add_page_break()
    para(doc, "von Post humification scale — field reference", bold=True, size=13, space_after=2)
    para(doc, "Squeeze a handful of WET peat and watch what comes out between your fingers.",
         size=9, italic=True, space_after=6)
    von_post_table(doc)
    doc.add_paragraph()
    para(doc, "H1–H3 = fibric peat (Oi)      ·      H4–H6 = hemic, “mucky peat” (Oe)      "
              "·      H7–H10 = sapric, “muck” (Oa)", bold=True, size=9, space_after=6)
    para(doc, "The more decomposed the material, the LESS carbon per unit weight — but also the "
              "HIGHER its bulk density, and those two partly cancel. Do not assume a muck layer "
              "holds less carbon per centimetre than a fibric one. Measure it.",
         size=8, italic=True, space_after=10)

    para(doc, "Three transitions worth a note every time", bold=True, size=11, space_after=2)
    tt = doc.add_table(rows=4, cols=3)
    borders(tt)
    fw = [Inches(0.8), Inches(1.5), Inches(4.8)]
    set_widths(tt, fw)
    rows = [("Frame", "Where", "The note that matters"),
            ("A", "0–50 cm, Sphagnum bog",
             "The depth where green, LIVING Sphagnum turns to brown peat. This is the "
             "vegetation/peat boundary."),
            ("B", "mid-core, 50–100 cm",
             "Dark mucky peat at the top grading to lighter peat below, and the depth of any "
             "significant WOOD."),
            ("C", "at the base",
             "The depth of the peat → mineral transition, the texture and colour of the mineral "
             "layer, and the depth where coarse fragments become common.")]
    for i, triple in enumerate(rows):
        for j, v in enumerate(triple):
            write(tt.rows[i].cells[j], v, bold=(i == 0 or j == 0), size=8)
            shade(tt.rows[i].cells[j], GREY if i == 0 else (TINT if i % 2 else WHITE))

    doc.add_paragraph()
    para(doc, "Sources: Measuring Carbon in Peat Soils: A Supplemental Guide (WWF-Canada, 2024). "
              "Core-length, surface-compaction, bore-depth and cylinder-diameter fields follow "
              "Bansal et al. (2023), Practical Guide to Measuring Wetland Carbon Pools and "
              "Fluxes, Wetlands 43:105, p.22.", size=7, italic=True, color="7F7F7F")

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path


# --------------------------------------------------------------------------- #
# 2. Peat Carbon Skill Checklist
# --------------------------------------------------------------------------- #
# ("SECTION", title) | ("SUB", title) | (skill, description, team?)
TEAM = "▶ TEAM"   # marks an item that cannot be signed off solo

CHECKLIST = [
    ("SECTION", "SECTION 1 — Sampling Design & Depth Survey"),
    ("SUB", "Study area and strata"),
    ("Define the study area", "Draw the boundary the estimate will apply to, and write down the RULE used "
     "(probed organic depth ≥ 30 cm / mapped wetland polygon / visible extent). Exclude open water.", False),
    ("Identify wetland type", "Bog, fen, swamp or marsh, from water source and vegetation. This sets the "
     "strata and whether the tree protocol applies.", False),
    ("Stratify within the wetland", "Split centre from margin, and note the depocentre. This is the single "
     "highest-return decision in the design — depth varies most between centre and edge.", False),
    ("Set minimum plot spacing", "At least the depth-grid spacing (10–25 m). Re-draw if randomisation puts "
     "two plots a few metres apart.", False),
    ("SUB", "Site-level depth survey"),
    ("Run the site depth grid", "Probe at every intersection of a 10–25 m grid across the site. Larger areas "
     "(≥ 10,000 ha) need mapping software rather than a walked grid.", False),
    ("Read a probe correctly", "Push until firm resistance, then read at the peat surface. Distinguish buried "
     "WOOD (abrupt, shifts if you move 20 cm) from the mineral contact (firm, gritty, consistent between "
     "neighbouring probes). Re-probe any reading that disagrees with its neighbours.", False),
    ("Confirm it is a peatland", "Organic depth ≥ 30 cm. If it is not, that is a finding — record and "
     "report it rather than quietly treating the site as peatland.", False),
    ("SUB", "How many cores"),
    ("Derive a variability prior", "Compute the CV of probed depth WITHIN a stratum. It predicts the CV of "
     "carbon stock to within about 10%. Never pool depth CV across wetland types — carbon density per cm "
     "differs between bog, fen and swamp.", False),
    ("Size the campaign", "Cochran's n for the design figure; field at least the t-corrected floor "
     "(about two more cores). Minimum 3 cores per stratum whatever the formula returns.", False),
    ("Decide on dating BEFORE the field", "Accumulation rates need 1–2 cm continuous sections through the "
     "upper 10–50 cm and ideally three replicate cores. This cannot be added afterwards.", False),

    ("SECTION", "SECTION 2 — Site Preparation & Documentation"),
    ("SUB", "Plot setup"),
    ("Lay out the plot", "10 × 10 m (100 m²) for peat. If tree cover is ≥ 25%, a 400 m² large plot "
     "around it for the tree survey.", False),
    ("Record coordinates at the CORING POINT", "Not the plot corner. Log latitude, longitude, elevation, "
     "datum AND GNSS accuracy — a handheld unit under a swamp canopy can be 5–10 m out.", False),
    ("Assign the CoreID", "A systematic unique identifier, location–site–sample (e.g. PE-01-B). Write the "
     "scheme on the datasheet header; mixed conventions are the commonest cause of data that will not join.", False),
    ("Record wetland type, microform, vegetation zone", "All three on the plot log. Microform is where the "
     "CORER goes in, not the plot average.", False),
    ("Measure water table depth", "Dig or probe a small hole beside the plot, let it equilibrate a few "
     "minutes, measure DOWN from the peat surface. Positive = below surface; negative = standing water above.", False),
    ("SUB", "Photo series — 14 photos per plot"),
    ("Photo — straight down", "Document the ground vegetation from the coring spot.", False),
    ("Photo — straight up", "Document canopy cover, or sky if there is none.", False),
    ("Photos — each cardinal direction", "For each of N, S, E, W: one parallel with the ground, one angled "
     "45° up, one angled 45° down. Twelve photos, fourteen in total.", False),
    ("SUB", "Order of work"),
    ("Complete ALL vegetation work before coring", "Coring is destructive. Probing is mildly destructive too "
     "— probe after the vegetation work, and keep your route through the plot deliberate.", False),
    ("Offset cores from permanent vegetation plots", "Otherwise the next visit measures regrowth on ground "
     "you dug up.", False),

    ("SECTION", "SECTION 3 — Coring"),
    ("SUB", "Within-plot depth survey and coring point"),
    ("Run the 1 m in-plot grid", "121 probe points in the 10 × 10 m plot, by digital grid or two tapes. "
     "Keep the depths — they justify the coring point and become your next prior.", False),
    ("Choose the coring point", "At the MEDIAN probed depth, free of obstructions. Move off it only for a "
     "physical reason, and write the reason down.", False),
    ("SUB", "Corer setup"),
    ("Identify and MARK the closed position", "Serrated edge on the inside edge of the guard = OPEN; in the "
     "middle of the guard = CLOSED. Mark closed with tape or paint before going to the field — you will be "
     "judging this by feel at the end of 3 m of rod.", False),
    ("Assemble the corer", "At least one extension rod, plus roughly one more per additional metre of depth. "
     "Handle on top.", False),
    ("Tape the target depth on the rods", "The only reliable stop. Pushing past the intended depth invalidates "
     "the sample — material from below is mixed in.", False),
    ("Prepare the processing area", "Flat spot near the coring point, tarp down, all equipment laid out "
     "before the first drive.", False),
    ("SUB", "Extracting a core — a three-to-four-person job"),
    ("Insert the corer OPEN and vertical", "A leaning corer overstates depth and shears the sample. Dense peat "
     "may need two or three people pushing, or a sledgehammer.", True),
    ("Turn the handle 180° to CLOSED", "One smooth rotation. Reversing part-way mixes the sample.", True),
    ("Extract under rotational pressure", "One person lifts from the handle keeping the barrel against the "
     "guard; a second lifts low on the rods, ready to catch and clasp barrel and guard together as it "
     "emerges. This is the moment most samples are lost.", True),
    ("Free a stuck corer safely", "Rock it gently to break the suction. Never lever it sideways — bent rods "
     "make depth unreliable and the corer unsafe.", True),
    ("Carry the core horizontal", "Sample resting on the guard, barrel facing up. A tilted core drains and "
     "slumps.", True),
    ("Work the drive sequence", "50 cm at a time — 0–50, 50–100, 100–150 — until refusal. Record the "
     "number of drives.", True),
    ("SUB", "Recording the core"),
    ("Record Bansal's four measurements", "Core length recovered; surface compaction (inserted-corer surface "
     "elevation minus true soil surface); bore depth; cylinder diameter.", False),
    ("Measure the chamber diameter with calipers", "It sets sample volume and therefore every bulk density. "
     "Out by 10% puts every carbon stock out by 21%. Write it on the corer.", False),
    ("Compute and check the recovery ratio", "Recovered length ÷ bore depth. Below 0.90 the chamber did not "
     "fill — material is MISSING. Never apply a push-corer compaction correction to a side-filling core.", False),
    ("Record whether the mineral contact was reached", "YES / NO, and the depth if yes. If no, the stock is a "
     "MINIMUM. Re-core 1–2 m away — buried wood is local, the contact is not.", False),
    ("Record the reason for refusal", "Buried wood, rock, gravel, rods bending, or the contact itself.", False),

    ("SECTION", "SECTION 4 — Reveal, Measure and Photograph"),
    ("SUB", "Revealing the core"),
    ("Reveal the core horizontally", "Corer on the tarp barrel-up; one person holds barrel and guard and "
     "opens it, a second turns the handle. Keep it horizontal throughout.", True),
    ("Measure the core length", "In centimetres, before anything is cut.", False),
    ("SUB", "Describing the core"),
    ("Describe colour and texture changes", "And whether each change is GRADUAL or DISTINCT — a sharp contact "
     "means something happened fast.", False),
    ("Note large objects and their depths", "Wood, rocks, roots, identifiable plant remains.", False),
    ("Note gaps and water saturation", "Missing areas where the barrel did not fill; mucky, semi-saturated "
     "or dry.", False),
    ("Assign von Post humification", "H1–H10, per section or at minimum at each distinct change plus at the "
     "base. H8–H10 beside mineral material means a contact; H3 at refusal is probably buried wood.", False),
    ("Note the three key transitions", "Frame A: green living Sphagnum → brown peat. Frame B: mucky → lighter "
     "peat, and the depth of significant wood. Frame C: peat → mineral, with mineral texture and colour.", False),
    ("SUB", "Photographing the core"),
    ("Photograph the revealed core", "Site and core number on the whiteboard, beside the core. Shoot from "
     "directly above with the ENTIRE profile and the whiteboard in frame. No shadows across the core.", False),

    ("SECTION", "SECTION 5 — Packaging or Sectioning"),
    ("SUB", "Field sectioning"),
    ("Section one at a time, top down", "Cut and package each section before cutting the next. Never cut the "
     "whole core and then bag it.", False),
    ("Cut at distinct layers", "So each section is a homogeneous mass. Sections at least 1 cm and no more "
     "than 5 cm.", False),
    ("Transfer with a trowel to a pre-labelled bag", "Keep the knife where it cut. Label CoreID + section "
     "number + DEPTH INTERVAL — the interval, not just the number.", False),
    ("Two people, two jobs", "One sections, one labels. A single person doing both is how a 70-bag core ends "
     "up with two bags labelled “40–45”.", True),
    ("Wash and dry tools between sections", "Reduces contamination between depths.", False),
    ("SUB", "Whole-core packaging"),
    ("Label TOP and BOTTOM on pipe and board", "Both ends, both items, before the core goes near it. A core "
     "with its orientation lost is worthless, and an inverted core produces an age model that runs backwards.", False),
    ("Line and load the PVC pipe", "Aluminium foil then plastic wrap; pipe over the core; flip corer and pipe "
     "together so the core falls into the cradle.", True),
    ("Wrap, board and tape", "Plastic wrap then foil, poster board over the top, duct-taped, labelled with "
     "the CoreID.", False),
    ("SUB", "Storage"),
    ("Cooler, then freezer", "Horizontal, board up, packed so nothing moves in transit. Freeze as soon as "
     "possible — warm wet peat respires and carbon content shifts.", False),
    ("Clean the corer and all tools", "Before the next core.", False),

    ("SECTION", "SECTION 6 — If You Are Dating the Core"),
    ("SUB", "Sectioning for chronology"),
    ("Section at 1–2 cm through the upper core", "Through the top 10–50 cm. 5 cm sections smear the "
     "¹³⁷Cs peak away and it cannot be recovered.", False),
    ("Keep the sequence CONTINUOUS", "The CRS ²¹⁰Pb model requires the entire inventory. A gap invalidates "
     "every date below it.", False),
    ("Section finely deeper than you think you need", "You cannot tell in the field where ²¹⁰Pb reaches "
     "supported background.", False),
    ("Use a clean blade and avoid metal contamination", "Radionuclide and trace-element work needs extra care. "
     "Wash between every section.", False),
    ("Bag a ¹⁴C target from the base separately", "Terrestrial seeds, leaves or twigs. NOT roots (grew down "
     "from above — too young) and NOT aquatic macrofossils (reservoir effect — too old). Note the depth.", False),
    ("Consider a box corer for the surface", "A Russian corer handles loose fibrous acrotelm poorly. A "
     "Wardenaar box corer for the top 50 cm gives an uncompressed monolith. Note the overlap depth so the "
     "two cores can be stitched.", False),

    ("SECTION", "SECTION 7 — Lab Chain & Calculations"),
    ("SUB", "Lab preparation"),
    ("Contact the lab in advance", "Confirm handling, packaging and delivery requirements. Ask WHICH "
     "bulk-density basis they report, and whether their volume assumes a full or half cylinder — that is a "
     "real factor-of-two error.", False),
    ("Request the right analyses", "Dry bulk density, moisture content, LOI₅₅₀. CHN on a subset if you can "
     "afford it.", False),
    ("Calibrate the LOI → carbon factor locally", "The default 0.5 is not universal — published factors run "
     "0.21 to 0.58. Run CHN on a subset of your own samples and calibrate against LOI.", False),
    ("SUB", "Calculations"),
    ("Section carbon stock", "Bulk density × (carbon % ÷ 100) × section thickness, converted to kg C/m².", False),
    ("Core stock, two bases", "Full profile to the mineral contact (the headline) AND to a fixed reference "
     "depth, default 100 cm (for comparability).", False),
    ("Plot and site means with an interval", "Mean, SD, standard error and a confidence interval using t, not "
     "z, at small sample size.", False),
    ("Check achieved against target precision", "Compare the relative margin of error to the target set in "
     "planning. If it is missed, scrutinise the raw data and post-stratify by landscape position before "
     "adding cores.", False),
    ("Scale to the study area", "Area-weighted across sites. State plainly that an area-weighted mean does "
     "not by itself propagate the per-site intervals.", False),
]


def build_checklist(path):
    doc = base_doc(margin=0.7)

    para(doc, "WWF-Canada Carbon Training", size=9, color="7F7F7F", space_after=0)
    para(doc, "Peat Carbon Sample Collection — In-Field Skill Checklist",
         bold=True, size=15, space_after=4)
    para(doc, "Participant Name: ________________________________    Date: _______________    "
              "Site: _______________", size=10, space_after=4)
    para(doc, "Instructions: work through each skill with your trainer. Check the left ☐ when you "
              "have practised the skill; check the right ☐ when your trainer confirms you can do it "
              "unsupervised. Items marked ▶ TEAM require three to four people and cannot be signed "
              "off alone — extracting a peat core is a team skill.",
         size=8, italic=True, space_after=8)

    nrows = 1 + len(CHECKLIST)
    t = doc.add_table(rows=nrows, cols=4)
    borders(t)
    widths = [Inches(0.35), Inches(1.5), Inches(4.4), Inches(0.85)]
    set_widths(t, widths)

    for j, h in enumerate(("", "Skill", "Description / Key Actions", "Complete ✔")):
        write(t.rows[0].cells[j], h, bold=True, size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 3) else None)
        shade(t.rows[0].cells[j], GREY)

    alt = 0
    for i, item in enumerate(CHECKLIST, start=1):
        if item[0] == "SECTION":
            band(t, i, item[1], ORANGE, color="FFFFFF", size=10)
            alt = 0
            continue
        if item[0] == "SUB":
            band(t, i, item[1], PALE, size=9)
            alt = 0
            continue
        skill, desc, team = item
        fill = TINT if alt % 2 else WHITE
        alt += 1
        row = t.rows[i]
        write(row.cells[0], "☐", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        write(row.cells[1], (TEAM + "  " if team else "") + skill, bold=True, size=8)
        write(row.cells[2], desc, size=8)
        write(row.cells[3], "☐", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        for j in range(4):
            shade(row.cells[j], fill)

    doc.add_paragraph()
    para(doc, "Trainer name: ________________________________     "
              "Signature: ________________________________     Date: _______________",
         size=10, space_after=8)
    para(doc, "Sources: Measuring Carbon in Peat Soils: A Supplemental Guide (WWF-Canada, 2024); "
              "Bansal et al. (2023), Practical Guide to Measuring Wetland Carbon Pools and Fluxes, "
              "Wetlands 43:105. Sampling-design items follow Part 2 of the Wetland Carbon Workshop.",
         size=7, italic=True, color="7F7F7F")

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path


if __name__ == "__main__":
    a = build_datasheet(os.path.join(WET, "03_Field_Methods", "datasheets", "Peat-Core-Data-Sheet.docx"))
    b = build_checklist(os.path.join(WET, "03_Field_Methods", "checklists", "Peat_Carbon_Skill_Checklist.docx"))
    for p in (a, b):
        print(f"wrote {os.path.relpath(p, os.path.dirname(WET))}  ({os.path.getsize(p):,} bytes)")
