"""Recalculate the built workbooks so they ship with CACHED VALUES.

openpyxl writes formulas but never their results. A workbook straight out of
the build scripts therefore reads back as all-None through any tool that does
not evaluate formulas itself -- GitHub's preview, most viewers, and the
verify_* scripts in this folder. Excel recalculates on open so a person opening
it sees numbers, but the file as committed holds none.

This pass loads each workbook in LibreOffice Calc and writes it back with the
results cached. Run it as the LAST build step:

    python3 build_grass_1.py
    python3 build_grass_2.py
    python3 build_grass_3.py
    python3 fill_grass_example.py
    python3 recalc_grass.py        # <- this
    python3 verify_grass_wex.py
    python3 verify_grass_p5.py

Needs `libreoffice-calc` (not just libreoffice-core):
    apt-get install -y libreoffice-calc
"""
import os, shutil, subprocess, sys, tempfile

GRASS = "/home/user/Terrestrial_Carbon_Workshops_V1/Grasslands"

TARGETS = [
    f"{GRASS}/04_Data_Interpretation/calculators/Grassland_Carbon_Calculator.xlsx",
    f"{GRASS}/Worked_Example/Grassland_Carbon_Calculator_WorkedExample.xlsx",
]


def recalc(path, profile):
    """Round-trip one workbook through LibreOffice, in place."""
    if not os.path.exists(path):
        print("  missing, skipped:", path)
        return False
    with tempfile.TemporaryDirectory() as tmp:
        # Convert in a scratch dir: soffice will not write over its own input.
        src = os.path.join(tmp, os.path.basename(path))
        shutil.copy2(path, src)
        out = os.path.join(tmp, "out")
        os.makedirs(out)
        r = subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}",
             "--headless", "--norestore",
             "--convert-to", "xlsx:Calc MS Excel 2007 XML",
             "--outdir", out, src],
            capture_output=True, text=True, timeout=900)
        produced = os.path.join(out, os.path.basename(path))
        if not os.path.exists(produced):
            print("  FAILED:", os.path.basename(path))
            print("   ", (r.stderr or r.stdout).strip()[:400])
            return False
        shutil.copy2(produced, path)
    return True


def cached_count(path):
    import openpyxl
    f = openpyxl.load_workbook(path)
    v = openpyxl.load_workbook(path, data_only=True)
    n = c = 0
    for sn in f.sheetnames:
        sf, sv = f[sn], v[sn]
        for row in sf.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    n += 1
                    if sv[cell.coordinate].value is not None:
                        c += 1
    return n, c


if __name__ == "__main__":
    ok = True
    with tempfile.TemporaryDirectory() as profile:
        for t in TARGETS:
            print(os.path.relpath(t, GRASS))
            if not recalc(t, profile):
                ok = False
                continue
            n, c = cached_count(t)
            # The blank calculator has nothing to compute from, so most of its
            # formulas legitimately cache as blank. Only the worked example
            # should come back substantially populated.
            print(f"  formulas {n}, cached {c}")
            if "WorkedExample" in t and c == 0:
                print("  ERROR: worked example still has no cached values.")
                ok = False
    sys.exit(0 if ok else 1)
