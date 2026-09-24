"""Assert every number in Grasslands/Worked_Example/README.md against the
RECALCULATED workbook.  Reads cached values, so run LibreOffice recalc first."""
import openpyxl, collections, statistics, math
from decimal import Decimal, ROUND_HALF_UP
def f(x, nd):
    """format with ROUND_HALF_UP -- Python's f-string uses banker's rounding,
    which disagrees with Excel (and with people) on exact .xxx5 values."""
    q=Decimal(1).scaleb(-nd)
    return str(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP))
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
GRASS = os.path.dirname(HERE)          # .../Grasslands
sys.path.insert(0, HERE)               # tdist.py lives beside this script

# The workbook must be RECALCULATED first -- openpyxl writes formulas without
# cached values, so a freshly built workbook reads back as all None.  Run
#   python3 recalc_grass.py
# as the last build step; it does the LibreOffice pass in place, so the file
# committed to the repo already carries its results.  (GRASS_WEX_XLSX still
# overrides the path if you want to check a copy.)
WB = os.environ.get('GRASS_WEX_XLSX') or os.path.join(
        GRASS, 'Worked_Example', 'Grassland_Carbon_Calculator_WorkedExample.xlsx')
MD = os.path.join(GRASS, 'Worked_Example', 'README.md')
wb=openpyxl.load_workbook(WB,data_only=True)
txt=open(MD,encoding='utf-8').read()
ok=[]
def need(s,label=None):
    good=s in txt; ok.append(good)
    print(('  ok   ' if good else ' FAIL ')+f'prose has {label or s!r}')
def chk(label,got,want,tol=0):
    good=abs(got-want)<=tol; ok.append(good)
    print(('  ok   ' if good else ' FAIL ')+f'{label}: {got} vs {want}')

ps=wb['5. Plot Summary']; plots=[]
for r in range(5,20):
    if not ps.cell(r,1).value: continue
    plots.append(dict(id=ps.cell(r,1).value,site=ps.cell(r,2).value,inc=ps.cell(r,3).value,
        s30=ps.cell(r,4).value,sfull=ps.cell(r,5).value,deep=ps.cell(r,6).value,
        root=ps.cell(r,7).value,veg=ps.cell(r,8).value))
chk('plot count',len(plots),9)

print('\n-- per-plot table --')
for p in plots:
    need(f"| `{p['id']}` | {p['site']} | {f(p['s30'],3)} |",f"{p['id']} soil to 30")
    need(f"| {f(p['root'],3)} | {f(p['veg'],3)} |",f"{p['id']} root/veg")

print('\n-- counts --')
sd=wb['2. Soil Data']; rt=wb['3. Root Biomass']; vg=wb['4. Vegetation Data']
ninc=sum(1 for r in range(5,90) if sd.cell(r,1).value)
nrt =sum(1 for r in range(5,90) if rt.cell(r,1).value)
nvg =sum(1 for r in range(5,30) if vg.cell(r,1).value)
chk('soil increments',ninc,34); chk('root fractions',nrt,68); chk('veg quadrats',nvg,9)
need('**34 soil\nincrements**','34 soil increments'); need('**68 root fractions**'); need('**9 clip-and-weigh quadrats**')

print('\n-- site summary --')
ss=wb['6. Site Summary']
for r,lab in [(5,'S1'),(6,'S2'),(7,'S3')]:
    area=ss.cell(r,2).value; sm=ss.cell(r,4).value; shw=ss.cell(r,6).value
    rm=ss.cell(r,10).value; rhw=ss.cell(r,12).value
    sp=ss.cell(r,9).value; rp=ss.cell(r,15).value
    need(f'| **{f(sm,2)} ±{f(shw,2)}** |',f'{lab} soil mean/hw')
    need(f'| {f(rm,3)} ±{f(rhw,3)} |',f'{lab} root mean/hw')
    pct=round(100*ss.cell(r,7).value)
    assert f'±{pct}%' in sp, sp
    need(f'±{pct}%',f'{lab} soil achieved')
    chk(f'{lab} area ha',area/10000,{'S1':64.0,'S2':9.5,'S3':21.0}[lab],0.001)
need('| **S1** grazed prairie | 3 |'); need('64.0 ha'); need('9.5 ha'); need('21.0 ha')

print('\n-- study area --')
tot=ss.cell(29,2).value; mean=ss.cell(30,2).value; co2=ss.cell(31,2).value; ar=ss.cell(28,2).value
chk('area ha',ar/10000,94.5,0.001); need('**94.5 ha**')
need(f'**{tot/1000:,.1f} t C**','total t C')
# The study-area interval, from the STUDY-AREA INTERVAL block below the totals.
se=ss.cell(34,2).value; dfree=ss.cell(35,2).value; half=ss.cell(36,2).value
ach=ss.cell(37,2).value; verdict=ss.cell(38,2).value
need(f'**{f(mean,3)} ± {f(half,3)} kg C/m²**','area-weighted mean with interval')
need(f'**±{round(100*ach):.0f}% at 90% confidence — MET**','study-area precision')
need(f'*({int(dfree)} design df)*','design df')
chk('stratified SE',round(se,4),0.3743,0.0002)
chk('study-area df',int(dfree),6,0)
chk('achieved margin <= soil target',1 if ach<=0.20 else 0,1,0)
ok.append(verdict.startswith('MET'))
print(('  ok   ' if ok[-1] else ' FAIL ')+f'workbook verdict: {verdict}')
need(f'**{co2:,.1f} t CO₂e**','t CO2e')
s1=ss.cell(5,16).value; s2=ss.cell(6,16).value; s3=ss.cell(7,16).value
need(f'S1 {s1/1000:,.1f} · S2 {s2/1000:,.1f} · S3 {s3/1000:,.1f}','per-site t C')
chk('site totals sum to study total',round((s1+s2+s3)/1000,1),round(tot/1000,1),0.05)

print('\n-- point 1: pool shares --')
sites=collections.defaultdict(lambda: collections.defaultdict(list))
for p in plots:
    for k in ('s30','root','veg'): sites[p['site']][k].append(p[k])
for s in ('S1','S2','S3'):
    so=statistics.mean(sites[s]['s30']); ro=statistics.mean(sites[s]['root']); ve=statistics.mean(sites[s]['veg'])
    need(f'| {f(so,3)} | {f(ro,3)} | {f(ve,3)} | **{f(100*ro/(so+ro),1)}%** | **{f(100*ve/(so+ro+ve),1)}%** |',f'{s} shares')
sh=[100*p['root']/(p['s30']+p['root']) for p in plots]
need(f'**{f(min(sh),1)}% to {f(max(sh),1)}%**','per-plot root share range')

print('\n-- point 2: ash correction --')
raw=sum(rt.cell(r,7).value for r in range(5,90) if rt.cell(r,1).value)
cor=sum(rt.cell(r,10).value for r in range(5,90) if rt.cell(r,1).value)
need(f'**{f(raw,3)} g to {f(cor,3)} g — a {f(100*(raw-cor)/raw,1)}% reduction.**','ash correction')

print('\n-- point 3: fine vs coarse --')
byc=collections.Counter()
for r in range(5,90):
    if rt.cell(r,1).value: byc[rt.cell(r,5).value]+=rt.cell(r,14).value
t=sum(byc.values())
for k,want in [('<=2 mm (fine)','≤2 mm (fine)'),('>2 mm (coarse)','>2 mm (coarse)')]:
    need(f'| {f(byc[k],3)} | **{f(100*byc[k]/t,1)}%** |' if 'coarse' not in k
         else f'| {f(byc[k],3)} | {f(100*byc[k]/t,1)}% |', want)

print('\n-- point 4: roots by depth --')
bd=collections.defaultdict(float)
for r in range(5,90):
    if rt.cell(r,1).value: bd[(rt.cell(r,3).value,rt.cell(r,4).value)]+=rt.cell(r,14).value
n=len(plots); cum=0
for k in sorted(bd):
    m=bd[k]/n; cum+=m
    if k==(20,22): continue           # BOS-03 truncated increment, folded into 20-30 in the prose
    need(f'| {f(m,3)} | {f(m/(k[1]-k[0]),4)} | {f(cum,3)} |' if k!=(30,60)
         else f'| {f(m,3)} | **{f(m/(k[1]-k[0]),4)}** | **{f(cum,3)}** |', f'{k[0]}-{k[1]} cm')
a=sum(bd[k] for k in bd if k[1]<=30)/n; b=sum(bd[k] for k in bd if k[0]>=30)/n
need(f'**{f(100*b/(a+b),1)}% of the root carbon','share below 30 cm')
d0=bd[(0,10)]/n/10; d3=bd[(30,60)]/n/30
need(f'**{f(100*d3/d0,0)}% of what it is at the','30-60 density vs surface')

print('\n-- point 4b: full profile vs 30 cm --')
deep=[p for p in plots if p['deep']==60]
chk('plots reaching 60 cm',len(deep),7)
rat=statistics.mean(p['sfull']/p['s30'] for p in deep)
need(f'**{f(100*(rat-1),0)}% more**','full-profile uplift'); need('seven plots that got there')
need(f'*another {f(100*(rat-1),0)}%* again','study-area caveat uplift')

print('\n-- point 5: tree threshold --')
s3=[p for p in plots if p['site']=='S3']
tree={'BOS-01':1.94,'BOS-02':2.31,'BOS-03':1.42}; cover={'BOS-01':28,'BOS-02':31,'BOS-03':22}
for r in range(5,20):
    pid=vg.cell(r,1).value
    if pid in tree:
        chk(f'{pid} tree C',vg.cell(r,14).value,tree[pid],1e-9)
        chk(f'{pid} cover',vg.cell(r,13).value,cover[pid],0)
m=statistics.mean(p['veg'] for p in s3)
# What the abandoned 25% canopy-cover rule would have cost: BOS-03 sits at 22%,
# so a strict cover gate drops its measured tree carbon. The workshop now counts
# any tree over 2 m, so 'as recorded' is the correct column.
strict=statistics.mean(p['veg']-(tree[p['id']] if cover[p['id']]<25 else 0) for p in s3)
need(f'| **{f(m,3)}** |','S3 veg as recorded'); need(f'| **{f(strict,3)}** |','S3 veg strict')
need(f'**{f(m-strict,3)} — {f(100*(m-strict)/m,1)}% of the pool**','threshold delta')
need('**22%, 28% and 31%**')

print('\n-- point 6: CVs --')
for s,want in [('S1',0.05),('S2',0.15),('S3',0.16)]:
    cv=statistics.stdev(sites[s]['s30'])/statistics.mean(sites[s]['s30'])
    chk(f'{s} soil CV',round(cv,2),want,0)
    cvr=statistics.stdev(sites[s]['root'])/statistics.mean(sites[s]['root'])
    print(f'      {s} root CV {cvr:.2f}')
need('**0.34, 0.34, 0.63**'); need('**0.05, 0.15, 0.16**')

print('\n-- the undetectable difference (shared with Part 5) --')
d=statistics.mean(sites['S2']['s30'])-statistics.mean(sites['S1']['s30'])
need(f'**+{f(d,2)} kg C/m²** (+{f(100*d/statistics.mean(sites["S1"]["s30"]),1)}%)','difference')
need('**−1.33 to +2.95 kg C/m²**'); need('takes **30**')

print('\n-- QC flags actually present in the workbook --')
allflags=' '.join(str(sd.cell(r,19).value) for r in range(5,90))+' '+ \
         ' '.join(str(ps.cell(r,11).value) for r in range(5,20))+' '+ \
         ' '.join(str(vg.cell(r,16).value) for r in range(5,20))
for frag in ['Bulk-density basis not confirmed','Cored shallower than the reporting depth',
             'the bottom of the core, not the bottom of the roots','ABOVE-ground only',
             'no further correction applied']:
    good=frag in allflags; ok.append(good)
    print(('  ok   ' if good else ' FAIL ')+f'workbook emits {frag!r}')
nbd=sum(1 for r in range(5,90) if sd.cell(r,19).value and 'not confirmed' in sd.cell(r,19).value)
chk('rows with unconfirmed BD basis',nbd,4); need('`GP-02`, 4 rows')
ncf=sum(1 for r in range(5,90) if sd.cell(r,19).value and 'no further correction' in sd.cell(r,19).value)
chk('informational CF rows',ncf,30); need('| 30 rows |')
print('\nall roots recorded Live (so the dead path is untested):',
      {rt.cell(r,6).value for r in range(5,90) if rt.cell(r,1).value})
need('every fraction here is recorded Live')

print(f'\n{sum(ok)}/{len(ok)} checks passed')
sys.exit(0 if all(ok) else 1)
