"""Fill the calculator with a constructed teaching dataset and save the worked example."""
import openpyxl, random, datetime
SRC = "/home/user/Terrestrial_Carbon_Workshops_V1/Forests/04_Data_Interpretation/calculators/Forest_Carbon_Calculator.xlsx"
OUT = "/home/user/Terrestrial_Carbon_Workshops_V1/Worked_Example_tmp.xlsx"
random.seed(20260911)
wb = openpyxl.load_workbook(SRC)

# ── 1. Plot & Site Log ───────────────────────────────────────────────────────
ws = wb["1. Plot & Site Log"]
plots = [
 # plot,  site, area, date,        location,           lat,      lon,     gnss, datum, elev, shape, L,  M,  S, sn, ew, allow
 ("MR-01","S1","MR","2026-07-14","Moose Ridge upland", 51.4412,-84.2018, 0.4,"WGS84",214,"circular",400,100,1, 3.2,-1.8,"Yes"),
 ("MR-02","S1","MR","2026-07-14","Moose Ridge upland", 51.4438,-84.1975, 0.5,"WGS84",219,"circular",400,100,1, 6.1, 2.4,"Yes"),
 ("MR-03","S1","MR","2026-07-15","Moose Ridge upland", 51.4401,-84.1932, 0.4,"WGS84",208,"circular",400,100,1,11.5, 4.0,"No"),
 ("BC-01","S2","MR","2026-07-16","Beaver Creek lowland",51.4288,-84.2211,0.6,"WGS84",186,"square", 400,100,1, 1.1,-0.7,"Yes"),
 ("BC-02","S2","MR","2026-07-16","Beaver Creek lowland",51.4265,-84.2168,0.5,"WGS84",184,"square", 400,100,1, 2.0, 1.2,"Yes"),
 ("BC-03","S2","MR","2026-07-17","Beaver Creek lowland",51.4301,-84.2245,0.7,"WGS84",188,"square", 400,100,1, 0.8,-0.4,"Yes"),
]
for i, p in enumerate(plots):
    for j, v in enumerate(p, start=1):
        ws.cell(5+i, j, v)
    ws.cell(5+i, 19, "")

# ── 3. Tree Data ─────────────────────────────────────────────────────────────
UP = [("Trembling aspen",0.28),("White birch",0.18),("Balsam fir",0.16),
      ("White spruce",0.16),("Jack pine",0.12),("Black spruce",0.10)]
LO = [("Black spruce",0.52),("Tamarack larch",0.20),("Balsam fir",0.16),("White birch",0.12)]
def pick(mix):
    x = random.random(); c = 0
    for s, w in mix:
        c += w
        if x <= c: return s
    return mix[-1][0]

ws = wb["3. Tree Data"]; r = 5
for pid, site, *_ in plots:
    upland = site == "S1"
    n = random.randint(11, 16) if upland else random.randint(14, 20)
    for k in range(1, n+1):
        sp = pick(UP if upland else LO)
        # lowland conifer stand: smaller stems
        dbh = round(random.lognormvariate(2.95 if upland else 2.55, 0.42), 1)
        dbh = min(dbh, 68.0)
        # height on a subset only — that is realistic field practice
        h = ""
        if k % 3 == 1:
            h = round(max(2.5, 1.3 + (26 if upland else 17) * (dbh/(dbh+9.5))), 1)
        for j, v in enumerate([pid, k, sp, dbh, h], start=1):
            ws.cell(r, j, v)
        r += 1
N_TREE_ROWS = r - 5

# ── 4. Understory Data ───────────────────────────────────────────────────────
ws = wb["4. Understory Data"]; r = 5
SHRUBS = ["Alder","Bog birch","Willow","Soap berry","Shrubby cinquefoil"]
for pid, site, *_ in plots:
    for k in range(1, random.randint(4, 7)):
        sp = random.choice(SHRUBS)
        L = round(random.uniform(0.6, 2.1), 2); W = round(L*random.uniform(0.6,1.1), 2)
        H = round(random.uniform(0.6, 1.9), 2)
        for j, v in enumerate([pid, f"M{k}", "Medium", "shrub", sp, "", L, W, H, ""], start=1):
            ws.cell(r, j, v)
        r += 1
    # one short-statured tree per plot
    ws.cell(r,1,pid); ws.cell(r,2,"M9"); ws.cell(r,3,"Medium"); ws.cell(r,4,"tree")
    ws.cell(r,5,"Short statured tree"); ws.cell(r,6,round(random.uniform(1.4,3.6),1)); r += 1
    # two clip quadrats per plot
    for k in (1, 2):
        ws.cell(r,1,pid); ws.cell(r,2,f"G{k}"); ws.cell(r,3,"Small"); ws.cell(r,4,"ground")
        ws.cell(r,5,"all shrubs"); ws.cell(r,10, round(random.uniform(95, 340), 1)); r += 1
N_UNDER_ROWS = r - 5

# ── 2. Core Log + 5. Soil Data ───────────────────────────────────────────────
cl = wb["2. Core Log"]; sd = wb["5. Soil Data"]; cr, sr = 5, 5
for pid, site, *_ in plots:
    upland = site == "S1"
    ncores = 2 if pid in ("MR-01", "BC-01") else 1
    for ci in range(1, ncores+1):
        cid = f"{pid}-C{ci}"
        # MR-03 is a shallow stony plot that stops short of 30 cm — a deliberate teaching case
        target = 22 if pid == "MR-03" else random.choice([30, 40, 50])
        hole = target
        core_len = round(target / random.uniform(1.00, 1.14), 1)
        for j, v in enumerate([cid, pid, round(51.43+random.uniform(-.01,.01),5),
                               round(-84.21+random.uniform(-.01,.01),5), "Soil core",
                               hole, core_len], start=1):
            cl.cell(cr, j, v)
        cl.cell(cr, 9, target); cl.cell(cr, 10, "Yes")
        cr += 1
        # slices
        edges = [0,5,10,15,20,30,40,50]
        edges = [e for e in edges if e <= target]
        if edges[-1] < target: edges.append(target)
        for si in range(len(edges)-1):
            top, bot = edges[si], edges[si+1]
            mid = (top+bot)/2
            cpct = round(max(0.9, (38 if upland else 47) * (0.88 ** mid) + random.uniform(-0.4, 1.6)), 2)
            bd   = round(min(1.55, (0.17 if upland else 0.13) + 0.021*mid + random.uniform(-0.02,0.03)), 3)
            cf   = 18 if pid == "MR-03" else random.choice([0, 0, 3, 6, 11])
            for j, v in enumerate([pid, cid, si+1, top, bot], start=1):
                sd.cell(sr, j, v)
            sd.cell(sr, 7, cf)
            sd.cell(sr, 8, "Fine earth / total volume")
            sd.cell(sr, 9, bd)
            sd.cell(sr, 11, cpct)
            sr += 1
N_CORE_ROWS, N_SOIL_ROWS = cr-5, sr-5

# ── 7. Site Summary: site IDs + areas ────────────────────────────────────────
ws = wb["7. Site Summary"]
ws.cell(5,1,"S1"); ws.cell(5,2,72000)
ws.cell(6,1,"S2"); ws.cell(6,2,48000)

wb.save(OUT)
print(f"trees={N_TREE_ROWS} understory={N_UNDER_ROWS} cores={N_CORE_ROWS} slices={N_SOIL_ROWS}")
print("saved", OUT)
