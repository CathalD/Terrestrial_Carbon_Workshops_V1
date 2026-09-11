import sys, json; sys.path.insert(0, "/tmp/claude-0/-home-user/af4fc136-c459-5930-bae9-67144880d936/scratchpad")
from build_calc_part1 import *
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

SCR = "/tmp/claude-0/-home-user/af4fc136-c459-5930-bae9-67144880d936/scratchpad"
Q = json.load(open(f"{SCR}/_rows.json"))
wb = openpyxl.load_workbook(f"{SCR}/_stage3.xlsx")
P1, C1, T1, U1, S1 = Q["P1"], Q["C1"], Q["T1"], Q["U1"], Q["S1"]
NP, NC, NT, NU, NS = Q["N_PLOT"], Q["N_CORE"], Q["N_TREE"], Q["N_UNDER"], Q["N_SOIL"]
PL, CL, TD, UD, SDD = "'1. Plot & Site Log'", "'2. Core Log'", "'3. Tree Data'", "'4. Understory Data'", "'5. Soil Data'"
def rg(sh, col, first, n): return f"{sh}!${col}${first}:${col}${first+n-1}"
PKEY, PSITE = rg(PL,'A',P1,NP), rg(PL,'B',P1,NP)
PLA, PMA, PSA, PEFF = rg(PL,'L',P1,NP), rg(PL,'M',P1,NP), rg(PL,'N',P1,NP), rg(PL,'R',P1,NP)

# ═════════════════════════════════════════ 6. PLOT SUMMARY
ws = wb.create_sheet("6. Plot Summary", 6)
ws["A1"] = "PLOT SUMMARY"; ws["A1"].font = Font(bold=True, size=14, color=NAVY)
ws["A2"] = ("One row per plot, calculated for you. Each pool is divided by ITS OWN plot area — trees by the large "
            "plot, shrubs by the medium plot, ground vegetation by the small plot — which is why they can be added "
            "together as kg C/m². Soil is the mean of the cores taken in that plot.")
ws["A2"].font = F_IT; ws.merge_cells("A2:R2"); ws["A2"].alignment = WRAP; ws.row_dimensions[2].height = 32
hdr = ["Plot ID", "Site ID", "Tree carbon (kg C)", "Large plot area (m²)", "Trees (kg C/m²)",
       "Shrub carbon (kg C)", "Medium plot area (m²)", "Shrubs (kg C/m²)",
       "Ground carbon (kg C)", "Small plot area (m²)", "Ground veg (kg C/m²)",
       "Understory total (kg C/m²)", "Cores in plot", "Soil to reporting depth (kg C/m²)",
       "Soil, full cores (kg C/m²)", "TOTAL (kg C/m²)", "QC flags", "_num", "_inc"]
header_row(ws, 4, hdr, widths=[13,10,13,13,12,13,13,12,13,13,12,13,10,14,14,13,44,9,7])
F = P1
f = {
 1: f'=IF(INDEX({PKEY},ROW()-{4})="","",INDEX({PKEY},ROW()-{4}))',
 2: f'=IF($A{{r}}="","",INDEX({PSITE},MATCH($A{{r}},{PKEY},0)))',
 3: f'=IF($A{{r}}="","",SUMIFS({rg(TD,"P",T1,NT)},{rg(TD,"A",T1,NT)},$A{{r}}))',
 4: f'=IF($A{{r}}="","",INDEX({PEFF},MATCH($A{{r}},{PKEY},0)))',
 5: '=IF(OR($A{r}="",$D{r}="",$D{r}=0),"",$C{r}/$D{r})',
 6: f'=IF($A{{r}}="","",SUMIFS({rg(UD,"O",U1,NU)},{rg(UD,"A",U1,NU)},$A{{r}},{rg(UD,"C",U1,NU)},"Medium"))',
 7: f'=IF($A{{r}}="","",INDEX({PMA},MATCH($A{{r}},{PKEY},0)))',
 8: '=IF(OR($A{r}="",$G{r}="",$G{r}=0),"",$F{r}/$G{r})',
 9: f'=IF($A{{r}}="","",SUMIFS({rg(UD,"O",U1,NU)},{rg(UD,"A",U1,NU)},$A{{r}},{rg(UD,"C",U1,NU)},"Small"))',
 10: f'=IF($A{{r}}="","",INDEX({PSA},MATCH($A{{r}},{PKEY},0)))',
 11: '=IF(OR($A{r}="",$J{r}="",$J{r}=0),"",$I{r}/$J{r})',
 12: '=IF($A{r}="","",IF($H{r}="",0,$H{r})+IF($K{r}="",0,$K{r}))',
 13: f'=IF($A{{r}}="","",COUNTIFS({rg(CL,"B",C1,NC)},$A{{r}}))',
 14: f'=IF(OR($A{{r}}="",$M{{r}}=0),"",IFERROR(AVERAGEIFS({rg(CL,"L",C1,NC)},{rg(CL,"B",C1,NC)},$A{{r}}),""))',
 15: f'=IF(OR($A{{r}}="",$M{{r}}=0),"",IFERROR(AVERAGEIFS({rg(CL,"M",C1,NC)},{rg(CL,"B",C1,NC)},$A{{r}}),""))',
 16: '=IF($A{r}="","",IF($E{r}="",0,$E{r})+IF($L{r}="",0,$L{r})+IF($N{r}="",0,$N{r}))',
 17: ('=IF($A{r}="","",'
      'IF($C{r}=0,"No tree data for this plot. ","")'
      '&IF($M{r}=0,"No soil cores for this plot — soil contributes zero to the total. ","")'
      '&IF(AND($F{r}=0,$I{r}=0),"No understory data — the total is trees + soil only. ","")'
      '&IF(OR($D{r}="",$D{r}=0),"No large plot area — tree carbon per m² cannot be calculated. ",""))'),
 18: '=IF(ISNUMBER($P{r}),$P{r},0)',
 19: '=IF(AND(ISNUMBER($P{r}),$A{r}<>""),1,0)',
}
for col, t in f.items():
    for r in range(5, 5 + NP):
        ws.cell(r, col, t.format(r=r))
for i, k in enumerate(["g"]*17 + ["g", "g"], start=1):
    paint(ws, 5, 4 + NP, i, "s" if i == 16 else "g")
ws.column_dimensions["R"].hidden = True; ws.column_dimensions["S"].hidden = True
ws.freeze_panes = "C5"
PS = "'6. Plot Summary'"
PS_SITE, PS_NUM, PS_INC = f"{PS}!$B$5:$B${4+NP}", f"{PS}!$R$5:$R${4+NP}", f"{PS}!$S$5:$S${4+NP}"

# ═════════════════════════════════════════ 7. SITE SUMMARY
ws = wb.create_sheet("7. Site Summary", 7)
ws["A1"] = "SITE SUMMARY"; ws["A1"].font = Font(bold=True, size=14, color=NAVY)
ws["A2"] = ("Type a Site ID and its area; everything else calculates. This is where the loop from Part 2 closes: "
            "the precision you asked for when you sized the campaign is checked against the precision you actually "
            "achieved. A point estimate with no interval cannot be checked, so the interval is not optional.")
ws["A2"].font = F_IT; ws.merge_cells("A2:N2"); ws["A2"].alignment = WRAP; ws.row_dimensions[2].height = 32
hdr = ["Site ID", "Site area (m²)", "Plots", "Mean (kg C/m²)", "SD", "Standard error",
       "t (two-sided)", "± half-width (kg C/m²)", "Achieved margin", "Target margin",
       "Precision", "Total carbon (kg C)", "Total (t CO₂e)", "QC flags"]
header_row(ws, 4, hdr, widths=[14, 14, 8, 14, 11, 12, 11, 15, 13, 12, 20, 15, 14, 40])
NSITE = 14
SR = 5
for i, k in enumerate(["y", "y"] + ["g"]*12, start=1):
    paint(ws, SR, SR+NSITE-1, i, k)
sf = {
 3: f'=IF($A{{r}}="","",SUMIFS({PS_INC},{PS_SITE},$A{{r}}))',
 4: f'=IF(OR($A{{r}}="",NOT(ISNUMBER($C{{r}})),$C{{r}}=0),"",SUMIFS({PS_NUM},{PS_SITE},$A{{r}})/$C{{r}})',
 5: (f'=IF(OR($A{{r}}="",NOT(ISNUMBER($C{{r}})),$C{{r}}<2),"",'
     f'SQRT(SUMPRODUCT(({PS_SITE}=$A{{r}})*{PS_INC}*({PS_NUM}-$D{{r}})^2)/($C{{r}}-1)))'),
 6: '=IF(OR(NOT(ISNUMBER($E{r})),NOT(ISNUMBER($C{r})),$C{r}<2),"",$E{r}/SQRT($C{r}))',
 7: '=IF(OR($A{r}="",NOT(ISNUMBER($C{r})),$C{r}<2),"",TINV(1-TARGET_CONFIDENCE,$C{r}-1))',
 8: '=IF(OR(NOT(ISNUMBER($F{r})),NOT(ISNUMBER($G{r}))),"",$G{r}*$F{r})',
 9: '=IF(OR(NOT(ISNUMBER($H{r})),NOT(ISNUMBER($D{r})),$D{r}=0),"",$H{r}/$D{r})',
 10: '=IF($A{r}="","",TARGET_MARGIN)',
 11: ('=IF(NOT(ISNUMBER($I{r})),IF($A{r}="","","Cannot be assessed — at least 2 plots are needed for an interval."),'
      'IF($I{r}<=$J{r},"MET: ±"&TEXT($I{r},"0%")&" at "&TEXT(TARGET_CONFIDENCE,"0%")&" confidence",'
      '"NOT MET: ±"&TEXT($I{r},"0%")&" against a ±"&TEXT($J{r},"0%")&" target"))'),
 12: '=IF(OR($A{r}="",NOT(ISNUMBER($B{r})),NOT(ISNUMBER($D{r}))),"",$D{r}*$B{r})',
 13: '=IF(NOT(ISNUMBER($L{r})),"",$L{r}*CO2E_FACTOR/1000)',
 14: ('=IF($A{r}="","",'
      'IF($C{r}=0,"No plots carry this Site ID — check it matches the Plot & Site Log. ","")'
      '&IF(AND($C{r}>0,$C{r}<2),"Only one plot: a mean can be reported but never an uncertainty. ","")'
      '&IF($B{r}="","No site area, so no total carbon. ","")'
      '&IF(AND(ISNUMBER($I{r}),$I{r}>$J{r}),"Wider than planned. More plots, or stratification, would tighten it — see Part 2, Appendix A8. ",""))'),
}
for col, t in sf.items():
    for r in range(SR, SR+NSITE):
        ws.cell(r, col, t.format(r=r))
# study-area block
B0 = SR + NSITE + 2
band(ws, B0, 14, "  STUDY AREA — all sites combined, weighted by area")
labs = [("Total area (m²)", f'=IF(SUM($B${SR}:$B${SR+NSITE-1})=0,"",SUM($B${SR}:$B${SR+NSITE-1}))'),
        ("Total carbon (kg C)", f'=IF(SUM($L${SR}:$L${SR+NSITE-1})=0,"",SUM($L${SR}:$L${SR+NSITE-1}))'),
        ("Area-weighted mean (kg C/m²)", f'=IF(OR($B${B0+1}="",$B${B0+2}=""),"",$B${B0+2}/$B${B0+1})'),
        ("Total (t CO₂e)", f'=IF($B${B0+2}="","",$B${B0+2}*CO2E_FACTOR/1000)')]
for i, (lab, fo) in enumerate(labs, start=1):
    ws.cell(B0+i, 1, lab).font = F_SUB
    c = ws.cell(B0+i, 2, fo); c.fill, c.border, c.font = FILL["s"], BOX, Font(bold=True, size=10)
ws.cell(B0+6, 1, "Reported to a soil depth of:").font = F_SUB
ws.cell(B0+6, 2, "=REPORTING_DEPTH_CM&\" cm\"").font = F_N
ws.cell(B0+7, 1, ("Note: the study-area mean above weights each site by its area, which is the right thing to do when "
                  "sites differ in size. It does NOT propagate the per-site uncertainties into a study-area interval — "
                  "that needs the stratified estimator described in Part 4.")).font = F_IT
ws.merge_cells(start_row=B0+7, start_column=1, end_row=B0+7, end_column=10)
ws.cell(B0+7, 1).alignment = WRAP; ws.row_dimensions[B0+7].height = 30
ws.freeze_panes = "C5"

wb.move_sheet("8. Settings", offset=-2)
wb.save(f"{SCR}/_stage4.xlsx")
print("sheets:", wb.sheetnames)
