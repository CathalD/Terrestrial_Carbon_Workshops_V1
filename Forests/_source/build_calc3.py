import sys, json; sys.path.insert(0, "/tmp/claude-0/-home-user/af4fc136-c459-5930-bae9-67144880d936/scratchpad")
from build_calc_part1 import *
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

SCR = "/tmp/claude-0/-home-user/af4fc136-c459-5930-bae9-67144880d936/scratchpad"
R = json.load(open(f"{SCR}/_refs.json"))
wb = openpyxl.load_workbook(f"{SCR}/_stage2.xlsx")
N_PLOT, N_CORE, N_TREE, N_UNDER, N_SOIL = 60, 120, 400, 250, 500
R1, R2 = "'R1. Tree Coefficients'", "'R2. Understory Coefficients'"
SPRANGE = f"{R1}!$L${R['SP_FIRST']}:$L${R['SP_LAST']}"
UNRANGE = f"{R2}!$I${R['UN_FIRST']}:$I${R['UN_LAST']}"


def build(title, pos, hdr, widths, fills, nrows, formulas, dvs=None, note=None):
    ws = wb.create_sheet(title, pos)
    ws["A1"] = title.split(". ", 1)[1].upper()
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    if note:
        ws["A2"] = note; ws["A2"].font = F_IT
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(hdr))
        ws["A2"].alignment = WRAP; ws.row_dimensions[2].height = 30
    HR = 4
    header_row(ws, HR, hdr, widths=widths)
    first = HR + 1
    for i, f in enumerate(fills, start=1):
        if f:
            paint(ws, first, first + nrows - 1, i, f)
    for col, tmpl in formulas.items():
        for r in range(first, first + nrows):
            ws.cell(r, col, tmpl.format(r=r))
    for dv_range, dv in (dvs or []):
        d = DataValidation(type="list", formula1=dv, allow_blank=True, showDropDown=False)
        ws.add_data_validation(d); d.add(dv_range.format(a=first, b=first + nrows - 1))
    ws.freeze_panes = ws.cell(first, 3)
    return ws, first


# ═════════════════════════════════════════ 1. PLOT & SITE LOG
hdr = ["Plot ID", "Site ID", "Study area", "Date", "Location", "Latitude", "Longitude",
       "GNSS accuracy (m)", "Datum", "Elevation (m)", "Plot shape",
       "Large plot area (m²)", "Medium plot area (m²)", "Small plot area (m²)",
       "Slope S–N (°)", "Slope E–W (°)", "Slope allowance applied in field?",
       "Effective large plot area (m²)", "Notes"]
w  = [14, 10, 12, 11, 18, 11, 11, 12, 10, 11, 12, 13, 13, 13, 11, 11, 15, 15, 30]
fl = ["y"] * 17 + ["g", "y"]
f  = {18: ('=IF($L{r}="","",IF($Q{r}="Yes",$L{r},'
           'IF(OR($O{r}="",$P{r}=""),$L{r},'
           '$L{r}*COS(RADIANS(ABS($O{r})))*COS(RADIANS(ABS($P{r}))))))')}
ws1, P1 = build("1. Plot & Site Log", 1, hdr, w, fl, N_PLOT, f,
  dvs=[("K{a}:K{b}", '"circular,square,rectangular"'),
       ("Q{a}:Q{b}", '"Yes,No"')],
  note="ONE ROW PER PLOT. This is the key every other sheet joins to. Fill this in first. "
       "Leave the plot-area columns at your standard sizes unless a plot was laid out differently. "
       "'Slope allowance applied in field?' — answer Yes if you lengthened the tape on the slope "
       "(Trees guide, Appendix); answer No if you laid out the nominal distance along the ground, "
       "and the effective area will be corrected for you.")
PLOTKEY = f"'1. Plot & Site Log'!$A${P1}:$A${P1+N_PLOT-1}"

# ═════════════════════════════════════════ 2. CORE LOG
hdr = ["Core ID", "Plot ID", "Latitude", "Longitude", "Sampling method",
       "Hole depth (cm)", "Core length recovered (cm)", "Compaction factor",
       "Depth to substratum (cm)", "Photo series taken?",
       "Deepest slice logged (cm)", "Stock to reporting depth (kg C/m²)",
       "Stock, full core (kg C/m²)", "Reaches reporting depth?", "QC flags", "Notes"]
w  = [14, 14, 11, 11, 15, 12, 13, 12, 13, 12, 13, 15, 15, 14, 40, 26]
fl = ["y"] * 7 + ["g"] + ["y", "y"] + ["g"] * 5 + ["y"]
SD = "'5. Soil Data'"
SDr = f"{SD}!$B$5:$B$%d" % (4 + N_SOIL)
f = {
 8:  '=IF(OR($F{r}="",$G{r}="",$G{r}=0),"",$F{r}/$G{r})',
 11: f'=IF($A{{r}}="","",IF(SUMPRODUCT(MAX(({SDr}=$A{{r}})*{SD}!$E$5:$E${4+N_SOIL}))=0,"",SUMPRODUCT(MAX(({SDr}=$A{{r}})*{SD}!$E$5:$E${4+N_SOIL}))))',
 12: f'=IF($A{{r}}="","",IFERROR(SUMIFS({SD}!$O$5:$O${4+N_SOIL},{SDr},$A{{r}}),""))',
 13: f'=IF($A{{r}}="","",IFERROR(SUMIFS({SD}!$P$5:$P${4+N_SOIL},{SDr},$A{{r}}),""))',
 14: '=IF($K{r}="","",IF($K{r}>=REPORTING_DEPTH_CM,"Yes","No"))',
 15: ('=IF($A{r}="","",'
      f'IF(AND($B{{r}}<>"",COUNTIF({PLOTKEY},$B{{r}})=0),"Plot ID not in Plot & Site Log. ","")'
      '&IF($K{r}="","No soil slices logged for this core. ","")'
      '&IF($N{r}="No","Core stops short of the reporting depth — its depth-limited stock is an underestimate. ","")'
      '&IF(AND($H{r}<>"",$H{r}>1.05),"Core compacted by "&TEXT(($H{r}-1)*100,"0")&"%: slice depths in the tube sit shallower than in the ground, so the depth-limited stock is biased. Full-core stock is unaffected. ","")'
      '&IF(AND($H{r}<>"",$H{r}<0.98),"Core length exceeds hole depth — check the two measurements. ",""))'),
}
ws2, C1 = build("2. Core Log", 2, hdr, w, fl, N_CORE, f,
  dvs=[("E{a}:E{b}", '"Soil core,Soil pit,Shallow soil"'),
       ("J{a}:J{b}", '"Yes,No"'),
       ("B{a}:B{b}", f"={PLOTKEY}")],
  note="ONE ROW PER CORE OR PIT. Matches the 'Core Notes' block on the paper Soil Carbon Data Sheet. "
       "Hole depth is the true depth of the ground you sampled; core length is what you recovered — the "
       "difference is compaction (Non-peat guide, Sampling).")
CORKEY = f"'2. Core Log'!$A${C1}:$A${C1+N_CORE-1}"

# ═════════════════════════════════════════ 3. TREE DATA
B = R["SP_FIRST"]
hdr = ["Plot ID", "Tree ID", "Species", "DBH (cm)", "Height (m)",
       "Coefficient block", "Equation used", "Wood (kg)", "Bark (kg)", "Branches (kg)",
       "Foliage (kg)", "Above-ground biomass (kg)", "Tree type", "Below-ground biomass (kg)",
       "Total biomass (kg)", "Carbon (kg C)", "QC flags"]
w  = [14, 9, 24, 10, 10, 11, 13, 11, 11, 11, 11, 14, 11, 14, 13, 12, 46]
fl = ["y", "y", "y", "y", "y"] + ["g"] * 12
def comp(off, outcol):
    return ('=IF(OR($D{r}="",$F{r}=""),"",'
            f'IF($E{{r}}>0,INDEX({R1}!$F:$F,$F{{r}}+{off})*($D{{r}}^INDEX({R1}!$G:$G,$F{{r}}+{off}))*($E{{r}}^INDEX({R1}!$H:$H,$F{{r}}+{off})),'
            f'INDEX({R1}!$D:$D,$F{{r}}+{off})*($D{{r}}^INDEX({R1}!$E:$E,$F{{r}}+{off}))))')
f = {
 6:  f'=IF($C{{r}}="","",IFERROR({B}+(MATCH($C{{r}},{SPRANGE},0)-1)*4,""))',
 7:  '=IF($F{r}="","",IF(AND($E{r}<>"",$E{r}>0),"DBH + height","DBH only"))',
 8:  comp(0, 8), 9: comp(1, 9), 10: comp(2, 10), 11: comp(3, 11),
 12: '=IF($F{r}="","",IFERROR(SUM($H{r}:$K{r}),""))',
 13: f'=IF($F{{r}}="","",INDEX({R1}!$J:$J,$F{{r}}))',
 14: ('=IF($L{r}="","",IF($M{r}="Deciduous",BGB_RS_DECIDUOUS_a*($L{r}^BGB_RS_DECIDUOUS_b),'
      'IF($M{r}="Conifers",BGB_RS_CONIFER*$L{r},'
      'AVERAGE(BGB_RS_DECIDUOUS_a*($L{r}^BGB_RS_DECIDUOUS_b),BGB_RS_CONIFER*$L{r}))))'),
 15: '=IF($L{r}="","",$L{r}+$N{r})',
 16: '=IF($O{r}="","",$O{r}*CARBON_FRACTION_BIOMASS)',
 17: ('=IF(AND($A{r}="",$C{r}=""),"",'
      f'IF(AND($A{{r}}<>"",COUNTIF({PLOTKEY},$A{{r}})=0),"Plot ID not in Plot & Site Log — this tree will not reach the summary. ","")'
      '&IF($A{r}="","No Plot ID. ","")'
      '&IF(AND($C{r}<>"",$F{r}=""),"Species not in the coefficient list — check the spelling against R1. ","")'
      '&IF($D{r}="","No DBH — nothing can be calculated. ","")'
      '&IF(AND($D{r}<>"",$D{r}>QC_DBH_MAX_CM),"DBH above "&QC_DBH_MAX_CM&" cm — check for a transcription error. ","")'
      '&IF($M{r}="All","Generic species: no published root:shoot applies, so below-ground biomass is the mean of the deciduous and conifer relationships. Choose a species, or Deciduous/Conifers, if you can. ","")'
      '&IF(AND($O{r}<>"",$O{r}>0,OR($N{r}/$O{r}<BGB_FRACTION_MIN,$N{r}/$O{r}>BGB_FRACTION_MAX)),"Below-ground share is "&TEXT($N{r}/$O{r},"0%")&", outside the 18–30% the Trees guide describes. ",""))'),
}
ws3, T1 = build("3. Tree Data", 3, hdr, w, fl, N_TREE, f,
  dvs=[("C{a}:C{b}", f"={SPRANGE}"), ("A{a}:A{b}", f"={PLOTKEY}")],
  note="ONE ROW PER TREE over 2 m tall, inside the large plot. Type Plot ID, Tree ID, species and DBH. "
       "Height is optional — leave it blank and the DBH-only equation is used; fill it in and the more "
       "precise DBH+height equation is used instead. Column G tells you which one ran.")
for c in (6,): ws3.column_dimensions[get_column_letter(c)].hidden = False

# ═════════════════════════════════════════ 4. UNDERSTORY DATA
hdr = ["Plot ID", "Sample ID", "Plot type", "Plant type", "Species / group",
       "Diameter at 0.3 m (cm)", "Length (m)", "Width (m)", "Height (m)",
       "Oven-dry clipped mass (g)", "Model parameter x", "Coefficient b", "Coefficient a",
       "Biomass (kg)", "Carbon (kg C)", "QC flags"]
w  = [14, 10, 11, 11, 22, 12, 10, 10, 10, 14, 12, 11, 11, 12, 12, 46]
fl = ["y"] * 9 + ["b"] + ["g"] * 6
f = {
 11: ('=IF($C{r}="Small","",IF($D{r}="tree",IF($F{r}="","",$F{r}),'
      'IF(OR($G{r}="",$H{r}="",$I{r}=""),"",$G{r}*$H{r}*$I{r})))'),
 12: f'=IF($C{{r}}="Small","",IFERROR(INDEX({R2}!$E${R["UC_FIRST"]}:$E${R["UC_LAST"]},MATCH($E{{r}},{UNRANGE},0)),""))',
 13: f'=IF($C{{r}}="Small","",IFERROR(INDEX({R2}!$F${R["UC_FIRST"]}:$F${R["UC_LAST"]},MATCH($E{{r}},{UNRANGE},0)),""))',
 14: ('=IF($C{r}="Small",IF($J{r}="","",$J{r}/1000),'
      'IF(OR($K{r}="",$L{r}="",$M{r}=""),"",($L{r}*($K{r}^$M{r}))/1000))'),
 15: '=IF($N{r}="","",$N{r}*CARBON_FRACTION_BIOMASS)',
 16: ('=IF(AND($A{r}="",$E{r}="",$J{r}=""),"",'
      f'IF(AND($A{{r}}<>"",COUNTIF({PLOTKEY},$A{{r}})=0),"Plot ID not in Plot & Site Log. ","")'
      '&IF($C{r}="","No plot type — say whether this came from a Medium or a Small plot. ","")'
      '&IF(AND($C{r}="Medium",$D{r}="tree",$F{r}=""),"Short-statured tree with no diameter at 0.3 m. ","")'
      '&IF(AND($C{r}="Medium",$D{r}="shrub",OR($G{r}="",$H{r}="",$I{r}="")),"Shrub needs all three of length, width and height. ","")'
      '&IF(AND($C{r}="Medium",$D{r}="shrub",$F{r}<>""),"Diameter entered for a shrub — it is ignored. Crown volume is used instead. ","")'
      '&IF(AND($C{r}="Medium",$E{r}<>"",$L{r}=""),"Species not in the coefficient list — use \'all shrubs\' if you have no closer match. ","")'
      '&IF(AND($C{r}="Small",$J{r}=""),"Small-plot row with no clipped dry mass yet. ","")'
      '&IF($N{r}<>"","Above-ground only: no root biomass is included for understory. ",""))'),
}
ws4, U1 = build("4. Understory Data", 4, hdr, w, fl, N_UNDER, f,
  dvs=[("C{a}:C{b}", '"Medium,Small"'), ("D{a}:D{b}", '"tree,shrub,ground"'),
       ("E{a}:E{b}", f"={UNRANGE}"), ("A{a}:A{b}", f"={PLOTKEY}")],
  note="TWO ROUTES, chosen by 'Plot type'. MEDIUM plots (0.5–2 m plants) use the allometric route: a "
       "short-statured tree needs its stem diameter at 0.3 m; a shrub needs crown length × width × height. "
       "SMALL plots (under 0.5 m) use clip-and-weigh: enter the oven-dry mass and nothing else.")

# ═════════════════════════════════════════ 5. SOIL DATA
hdr = ["Plot ID", "Core ID", "Sample ID", "Top depth (cm)", "Bottom depth (cm)",
       "Thickness (cm)", "Coarse fragments (% vol)", "Bulk density basis",
       "Bulk density (g/cm³)", "LOI₅₅₀ (%)", "Organic carbon (%)",
       "Carbon % used", "Coarse-fragment factor", "Thickness within reporting depth (cm)",
       "Stock to reporting depth (kg C/m²)", "Stock, full thickness (kg C/m²)", "QC flags", "Notes"]
w  = [13, 13, 10, 11, 11, 11, 13, 20, 12, 11, 12, 11, 13, 15, 15, 15, 46, 26]
fl = ["y"] * 8 + ["b", "b", "b"] + ["g"] * 6 + ["y"]
f = {
 6:  '=IF(OR($D{r}="",$E{r}=""),"",$E{r}-$D{r})',
 12: '=IF($K{r}<>"",$K{r},IF($J{r}<>"",$J{r}*CARBON_FRACTION_OM,""))',
 13: '=IF($H{r}="Fine earth / fine-earth volume",1-IF($G{r}="",0,$G{r})/100,1)',
 14: ('=IF(OR($D{r}="",$E{r}=""),"",'
      'MAX(0,MIN($E{r},REPORTING_DEPTH_CM)-MIN($D{r},REPORTING_DEPTH_CM)))'),
 15: '=IF(OR($I{r}="",$L{r}="",$N{r}=""),"",$I{r}*($L{r}/100)*$N{r}*$M{r}*10)',
 16: '=IF(OR($I{r}="",$L{r}="",$F{r}=""),"",$I{r}*($L{r}/100)*$F{r}*$M{r}*10)',
 17: ('=IF(AND($A{r}="",$B{r}="",$D{r}=""),"",'
      f'IF(AND($B{{r}}<>"",COUNTIF({CORKEY},$B{{r}})=0),"Core ID not in Core Log. ","")'
      f'&IF(AND($A{{r}}<>"",COUNTIF({PLOTKEY},$A{{r}})=0),"Plot ID not in Plot & Site Log. ","")'
      '&IF(AND($D{r}<>"",$E{r}<>"",$E{r}<=$D{r}),"Bottom depth is not below top depth. ","")'
      '&IF($H{r}="","Say which bulk-density basis the lab reported, or the coarse-fragment correction cannot be applied safely. ","")'
      '&IF(AND($I{r}<>"",OR($I{r}<QC_BD_MIN,$I{r}>QC_BD_MAX)),"Bulk density outside "&QC_BD_MIN&"–"&QC_BD_MAX&" g/cm³. ","")'
      '&IF(AND($L{r}<>"",$L{r}>QC_CARBON_PCT_MAX),"Carbon above "&QC_CARBON_PCT_MAX&"% — check units; this may be organic matter rather than carbon. ","")'
      '&IF(AND($J{r}<>"",$K{r}<>""),"Both LOI and measured carbon given — the measured value is used. ","")'
      '&IF(AND($N{r}<>"",$N{r}=0),"Slice lies entirely below the reporting depth — it counts to the full-core stock only. ",""))'),
}
ws5, S1 = build("5. Soil Data", 5, hdr, w, fl, N_SOIL, f,
  dvs=[("H{a}:H{b}", '"Fine earth / total volume,Fine earth / fine-earth volume"'),
       ("B{a}:B{b}", f"={CORKEY}"), ("A{a}:A{b}", f"={PLOTKEY}")],
  note="ONE ROW PER SLICE or soil layer. Depths run continuously down the core, so the bottom of one slice "
       "is the top of the next. Enter either LOI₅₅₀ or a measured organic carbon % — if you have both, the "
       "measured value wins. READ THE BULK DENSITY BASIS NOTE on the Instructions sheet before filling column H: "
       "picking the wrong one double-corrects for stones.")

wb.save(f"{SCR}/_stage3.xlsx")
json.dump(dict(P1=P1, C1=C1, T1=T1, U1=U1, S1=S1, N_PLOT=N_PLOT, N_CORE=N_CORE,
               N_TREE=N_TREE, N_UNDER=N_UNDER, N_SOIL=N_SOIL), open(f"{SCR}/_rows.json", "w"))
print("data sheets built:", [s for s in wb.sheetnames])
